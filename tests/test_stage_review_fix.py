"""
Task 6: the review (task3) and fix (task4) stages in isolation, via the mocked crew.
Since the actual "apply the review feedback" reasoning happens inside the (mocked-out)
LLM call, these tests focus on what run_pipeline's own code guarantees: review comments
are reported faithfully, and whatever the fix stage returns flows through to the final
test cases/report without the pipeline itself dropping, merging, or silently altering
anything.
"""

import json

from openpyxl import load_workbook

import agent


def test_review_comments_are_reported_correctly(
    patch_crew_kickoff, make_stage_outputs, sample_review_comments, sample_requirements_text, tmp_output_dir
):
    patch_crew_kickoff(make_stage_outputs(review=sample_review_comments))
    result = agent.run_pipeline(sample_requirements_text, output_dir=tmp_output_dir)

    wb = load_workbook(result["files"]["review"])
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    assert rows[0] == ("ID", "Comment")
    assert rows[1:] == [(c["id"], c["review_comment"]) for c in sample_review_comments]


def test_all_fixed_test_case_ids_are_carried_through_with_none_dropped(
    patch_crew_kickoff, make_stage_outputs, sample_test_cases, sample_requirements_text, tmp_output_dir
):
    patch_crew_kickoff(make_stage_outputs(fix=sample_test_cases))
    result = agent.run_pipeline(sample_requirements_text, output_dir=tmp_output_dir)

    final_ids = [tc["id"] for tc in result["data"]["final_test_cases"]]
    expected_ids = [tc["id"] for tc in sample_test_cases]

    assert final_ids == expected_ids
    assert len(final_ids) == len(set(final_ids)), "no ids should be merged/deduplicated by the pipeline"
    assert result["counts"]["final_test_cases"] == len(sample_test_cases)


def test_final_cases_excel_contains_every_test_case_from_the_fix_stage(
    patch_crew_kickoff, make_stage_outputs, sample_test_cases, sample_requirements_text, tmp_output_dir
):
    patch_crew_kickoff(make_stage_outputs(fix=sample_test_cases))
    result = agent.run_pipeline(sample_requirements_text, output_dir=tmp_output_dir)

    wb = load_workbook(result["files"]["final"])
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))[1:]

    assert len(rows) == len(sample_test_cases)
    assert {row[0] for row in rows} == {tc["id"] for tc in sample_test_cases}


def test_fix_stage_scenario_text_wins_over_draft_scenario_text(
    patch_crew_kickoff, sample_requirements_text, tmp_output_dir
):
    # scenario_lookup (used by the automation feasibility report) is built from the
    # FIX stage's output, not the draft generate-stage output — verify that wiring.
    draft = [{"id": "TC_1", "req_id": "REQ_1", "scenario": "DRAFT scenario text", "type": "Positive",
              "steps": ["a"], "expected_result": "x", "priority": "High"}]
    fixed = [{"id": "TC_1", "req_id": "REQ_1", "scenario": "FIXED scenario text", "type": "Positive",
              "steps": ["a"], "expected_result": "x", "priority": "High"}]
    automation = [{"id": "TC_1", "automatable": "Yes", "recommended_tool": "Selenium", "reason": "ok"}]

    patch_crew_kickoff({
        "analyze": "[]",
        "generate": json.dumps(draft),
        "review": "[]",
        "fix": json.dumps(fixed),
        "automation": json.dumps(automation),
        "estimate": "{}",
    })

    result = agent.run_pipeline(sample_requirements_text, output_dir=tmp_output_dir)

    wb = load_workbook(result["files"]["automation"])
    ws = wb.active
    row = list(ws.iter_rows(values_only=True))[1]
    assert row[1] == "FIXED scenario text"


def test_fix_stage_can_legitimately_drop_a_test_case_flagged_invalid(
    patch_crew_kickoff, sample_requirements_text, tmp_output_dir
):
    # The pipeline itself doesn't enforce "never drop" — that instruction lives in the
    # LLM prompt. If the fix stage output legitimately omits an id, the pipeline just
    # reports what it was given; this documents that the plumbing is a faithful pass-through.
    draft = [
        {"id": "TC_1", "req_id": "REQ_1", "scenario": "Keep me", "type": "Positive",
         "steps": ["a"], "expected_result": "x", "priority": "High"},
        {"id": "TC_2", "req_id": "REQ_1", "scenario": "Flagged as an exact duplicate", "type": "Positive",
         "steps": ["a"], "expected_result": "x", "priority": "High"},
    ]
    fixed = [draft[0]]  # TC_2 removed, as if the reviewer flagged it as a duplicate

    patch_crew_kickoff({
        "analyze": "[]", "generate": json.dumps(draft),
        "review": "[]", "fix": json.dumps(fixed),
        "automation": "[]", "estimate": "{}",
    })

    result = agent.run_pipeline(sample_requirements_text, output_dir=tmp_output_dir)

    final_ids = [tc["id"] for tc in result["data"]["final_test_cases"]]
    assert final_ids == ["TC_1"]
    assert result["counts"]["generated_test_cases"] == 2
    assert result["counts"]["final_test_cases"] == 1
