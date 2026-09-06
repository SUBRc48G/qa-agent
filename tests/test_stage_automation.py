"""
Task 7: the automation-feasibility stage (task5 / "automation") in isolation, via the
mocked crew. Covers tool recommendations, the summary calculation (total/automatable/
coverage %), and tool_counts aggregation.
"""

import json

from openpyxl import load_workbook

import agent


def test_automation_summary_totals_and_coverage_percentage(
    patch_crew_kickoff, make_stage_outputs, sample_test_cases, sample_automation_results,
    sample_requirements_text, tmp_output_dir,
):
    patch_crew_kickoff(make_stage_outputs(fix=sample_test_cases, automation=sample_automation_results))
    result = agent.run_pipeline(sample_requirements_text, output_dir=tmp_output_dir)

    summary = result["automation_summary"]
    assert summary["total"] == 5
    assert summary["automatable"] == 4
    assert summary["manual"] == 1
    assert summary["coverage_pct"] == 80.0


def test_tool_counts_are_aggregated_correctly(
    patch_crew_kickoff, make_stage_outputs, sample_test_cases, sample_automation_results,
    sample_requirements_text, tmp_output_dir,
):
    patch_crew_kickoff(make_stage_outputs(fix=sample_test_cases, automation=sample_automation_results))
    result = agent.run_pipeline(sample_requirements_text, output_dir=tmp_output_dir)

    assert result["automation_summary"]["tool_counts"] == {
        "Selenium": 2, "Playwright": 1, "JMeter": 1,
    }


def test_manual_only_cases_are_excluded_from_tool_counts(
    patch_crew_kickoff, make_stage_outputs, sample_test_cases, sample_automation_results,
    sample_requirements_text, tmp_output_dir,
):
    patch_crew_kickoff(make_stage_outputs(fix=sample_test_cases, automation=sample_automation_results))
    result = agent.run_pipeline(sample_requirements_text, output_dir=tmp_output_dir)

    assert "Manual" not in result["automation_summary"]["tool_counts"]


def test_zero_percent_coverage_when_nothing_is_automatable(
    patch_crew_kickoff, sample_requirements_text, tmp_output_dir
):
    all_manual = [
        {"id": "TC_1", "automatable": "No", "recommended_tool": "Manual", "reason": "exploratory"},
        {"id": "TC_2", "automatable": "No", "recommended_tool": "Manual", "reason": "usability"},
    ]
    patch_crew_kickoff({
        "analyze": "[]", "generate": "[]", "review": "[]", "fix": "[]",
        "automation": json.dumps(all_manual), "estimate": "{}",
    })
    result = agent.run_pipeline(sample_requirements_text, output_dir=tmp_output_dir)

    summary = result["automation_summary"]
    assert summary["coverage_pct"] == 0
    assert summary["automatable"] == 0
    assert summary["tool_counts"] == {}


def test_full_percent_coverage_when_everything_is_automatable(
    patch_crew_kickoff, sample_requirements_text, tmp_output_dir
):
    all_automatable = [
        {"id": "TC_1", "automatable": "Yes", "recommended_tool": "Postman", "reason": "API test"},
        {"id": "TC_2", "automatable": "Yes", "recommended_tool": "Postman", "reason": "API test"},
    ]
    patch_crew_kickoff({
        "analyze": "[]", "generate": "[]", "review": "[]", "fix": "[]",
        "automation": json.dumps(all_automatable), "estimate": "{}",
    })
    result = agent.run_pipeline(sample_requirements_text, output_dir=tmp_output_dir)

    summary = result["automation_summary"]
    assert summary["coverage_pct"] == 100.0
    assert summary["tool_counts"] == {"Postman": 2}


def test_automatable_check_is_case_and_whitespace_insensitive(
    patch_crew_kickoff, sample_requirements_text, tmp_output_dir
):
    mixed_case = [
        {"id": "TC_1", "automatable": " yes ", "recommended_tool": "Selenium", "reason": "x"},
        {"id": "TC_2", "automatable": "YES", "recommended_tool": "Selenium", "reason": "x"},
        {"id": "TC_3", "automatable": "no", "recommended_tool": "Manual", "reason": "x"},
    ]
    patch_crew_kickoff({
        "analyze": "[]", "generate": "[]", "review": "[]", "fix": "[]",
        "automation": json.dumps(mixed_case), "estimate": "{}",
    })
    result = agent.run_pipeline(sample_requirements_text, output_dir=tmp_output_dir)

    assert result["automation_summary"]["automatable"] == 2
    assert result["automation_summary"]["manual"] == 1


def test_automation_feasibility_excel_has_data_and_summary_sheets(
    patch_crew_kickoff, make_stage_outputs, sample_test_cases, sample_automation_results,
    sample_requirements_text, tmp_output_dir,
):
    patch_crew_kickoff(make_stage_outputs(fix=sample_test_cases, automation=sample_automation_results))
    result = agent.run_pipeline(sample_requirements_text, output_dir=tmp_output_dir)

    wb = load_workbook(result["files"]["automation"])
    assert wb.sheetnames == ["Automation Feasibility", "Summary"]

    data_ws = wb["Automation Feasibility"]
    data_rows = list(data_ws.iter_rows(values_only=True))
    assert data_rows[0] == ("ID", "Scenario", "Automatable", "Recommended Tool", "Reason")
    assert len(data_rows) == 1 + len(sample_automation_results)

    summary_ws = wb["Summary"]
    summary_rows = list(summary_ws.iter_rows(values_only=True))
    assert ("Total Test Cases", 5) in summary_rows
    assert ("Automatable", 4) in summary_rows
    assert ("Manual Only", 1) in summary_rows
    assert ("Automation Coverage %", 80.0) in summary_rows
