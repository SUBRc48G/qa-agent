"""
Task 8: the estimation stage (task6 / "estimate") in isolation, via the mocked crew.
Covers the estimation JSON flowing through correctly, the Excel report, the fallback
when the LLM's JSON omits test_cases_per_day, and that a different testcases_per_day
argument actually reaches the estimation task's prompt.
"""

import json

from openpyxl import load_workbook

import agent


def test_estimation_flows_through_to_the_result_unchanged(
    patch_crew_kickoff, make_stage_outputs, sample_estimation, sample_requirements_text, tmp_output_dir
):
    patch_crew_kickoff(make_stage_outputs(estimate=sample_estimation))
    result = agent.run_pipeline(sample_requirements_text, output_dir=tmp_output_dir)

    assert result["estimation"] == sample_estimation


def test_estimation_excel_report_has_correct_numbers(
    patch_crew_kickoff, make_stage_outputs, sample_estimation, sample_requirements_text, tmp_output_dir
):
    patch_crew_kickoff(make_stage_outputs(estimate=sample_estimation))
    result = agent.run_pipeline(sample_requirements_text, output_dir=tmp_output_dir)

    wb = load_workbook(result["files"]["estimation"])
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    assert rows[0] == ("Total TC", "Per Day", "Days")
    assert rows[1] == (
        sample_estimation["total_test_cases"],
        sample_estimation["test_cases_per_day"],
        sample_estimation["estimated_days"],
    )


def test_missing_test_cases_per_day_in_llm_output_falls_back_to_the_argument(
    patch_crew_kickoff, sample_requirements_text, tmp_output_dir
):
    incomplete_estimation = {"total_test_cases": 12, "estimated_days": 3}  # no test_cases_per_day key
    patch_crew_kickoff({
        "analyze": "[]", "generate": "[]", "review": "[]", "fix": "[]",
        "automation": "[]", "estimate": json.dumps(incomplete_estimation),
    })

    result = agent.run_pipeline(sample_requirements_text, testcases_per_day=7, output_dir=tmp_output_dir)

    wb = load_workbook(result["files"]["estimation"])
    ws = wb.active
    row = list(ws.iter_rows(values_only=True))[1]
    assert row == (12, 7, 3)


def test_different_testcases_per_day_values_reach_the_estimation_prompt(
    patch_crew_kickoff, make_stage_outputs, sample_estimation, sample_requirements_text, tmp_output_dir
):
    for rate in (5, 20, 100):
        patch_crew_kickoff(make_stage_outputs(estimate=sample_estimation))
        agent.run_pipeline(sample_requirements_text, testcases_per_day=rate, output_dir=tmp_output_dir)

        estimate_task = next(t for t in patch_crew_kickoff.captured_tasks if t.name == "estimate")
        assert f'"test_cases_per_day": {rate}' in estimate_task.description


def test_default_testcases_per_day_is_used_when_not_specified(
    patch_crew_kickoff, make_stage_outputs, sample_estimation, sample_requirements_text, tmp_output_dir
):
    patch_crew_kickoff(make_stage_outputs(estimate=sample_estimation))
    agent.run_pipeline(sample_requirements_text, output_dir=tmp_output_dir)

    estimate_task = next(t for t in patch_crew_kickoff.captured_tasks if t.name == "estimate")
    assert f'"test_cases_per_day": {agent.default_testcases_per_day}' in estimate_task.description


def test_zero_estimated_days_is_still_written_to_the_report(
    patch_crew_kickoff, sample_requirements_text, tmp_output_dir
):
    zero_estimation = {"total_test_cases": 0, "test_cases_per_day": 20, "estimated_days": 0}
    patch_crew_kickoff({
        "analyze": "[]", "generate": "[]", "review": "[]", "fix": "[]",
        "automation": "[]", "estimate": json.dumps(zero_estimation),
    })

    result = agent.run_pipeline(sample_requirements_text, output_dir=tmp_output_dir)

    assert "estimation" in result["files"]
    wb = load_workbook(result["files"]["estimation"])
    row = list(wb.active.iter_rows(values_only=True))[1]
    assert row == (0, 20, 0)
