"""
Shared pytest fixtures for the QA agent test suite.

Design note: most tests never call the real OpenAI API. Instead, `patch_crew_kickoff`
monkeypatches `crewai.Crew.kickoff` so it fakes each task's `.output` based on the task's
`name` (analyze/generate/review/fix/automation/estimate — see agent.STAGE_LABELS), while
the rest of `run_pipeline` (guardrail wiring, JSON extraction, warnings, Excel/docx writing)
runs for real. Tests that need the real API are marked `@pytest.mark.integration` and are
skipped by default (see test_pipeline_integration.py).
"""

import json
import os
import sys
from pathlib import Path

import pytest
from openpyxl import Workbook
from docx import Document as DocxDocument
from crewai import Crew
from crewai.tasks.task_output import TaskOutput

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import agent  # noqa: E402  (import after sys.path setup)


class TestConfig:
    """Constants shared across the test suite."""

    PROJECT_ROOT = Path(__file__).resolve().parent.parent
    API_KEY = os.getenv("OPENAI_API_KEY", "test-api-key")
    SAMPLE_PDF_PATH = PROJECT_ROOT / "requirements.pdf"
    SAMPLE_TXT_PATH = PROJECT_ROOT / "qa_requirements.txt"
    DEFAULT_TESTCASES_PER_DAY = 20
    STAGE_NAMES = ["analyze", "generate", "review", "fix", "automation", "estimate"]


@pytest.fixture
def test_config():
    return TestConfig


# ---------------- OUTPUT DIRECTORIES ----------------

@pytest.fixture
def tmp_output_dir(tmp_path):
    """A fresh, empty directory for a single run_pipeline() call to write reports into."""
    out = tmp_path / "outputs"
    out.mkdir()
    return str(out)


# ---------------- SAMPLE REQUIREMENTS TEXT ----------------

@pytest.fixture
def sample_requirements_text():
    return (
        "As a Network Account Read-only User, I want to list all the credentials "
        "created in the system, so that I can review the existing credentials.\n"
        "As a Network Account Read-Write User, I want to delete credentials from scope.\n"
        "As a Product Manager, I want CNUM operations to be enabled only on those NEs "
        "for which the relevant license is available.\n"
    )


# ---------------- MOCK DATA: REQUIREMENTS ----------------

@pytest.fixture
def sample_requirement_items():
    return [
        {"req_id": "REQ_1", "requirement": "List all credentials in the system", "category": "Functional", "priority": "High"},
        {"req_id": "REQ_2", "requirement": "Delete credentials from scope", "category": "Functional", "priority": "High"},
        {"req_id": "REQ_3", "requirement": "Enable CNUM operations only when licensed", "category": "Non-Functional", "priority": "Medium"},
    ]


# ---------------- MOCK DATA: TEST CASES ----------------

@pytest.fixture
def sample_test_cases():
    return [
        {"id": "TC_1", "req_id": "REQ_1", "scenario": "List credentials with valid session", "type": "Positive",
         "steps": ["Log in as Read-only User", "Navigate to credentials list", "Verify list is displayed"],
         "expected_result": "All credentials are listed", "priority": "High"},
        {"id": "TC_2", "req_id": "REQ_1", "scenario": "List credentials with no session", "type": "Negative",
         "steps": ["Attempt to open credentials list without logging in"],
         "expected_result": "User is redirected to login", "priority": "High"},
        {"id": "TC_3", "req_id": "REQ_1", "scenario": "List credentials when none exist", "type": "Edge",
         "steps": ["Log in", "Navigate to credentials list with zero credentials configured"],
         "expected_result": "An empty-state message is shown", "priority": "Medium"},
        {"id": "TC_4", "req_id": "REQ_2", "scenario": "Read-only user attempts delete", "type": "Security",
         "steps": ["Log in as Read-only User", "Attempt to delete a credential"],
         "expected_result": "Action is denied", "priority": "High"},
        {"id": "TC_5", "req_id": "REQ_1", "scenario": "List credentials with 10,000 records", "type": "Performance",
         "steps": ["Seed 10,000 credentials", "Log in", "Load the credentials list", "Measure response time"],
         "expected_result": "List loads within acceptable response time", "priority": "Medium"},
    ]


@pytest.fixture
def sample_review_comments():
    return [
        {"id": "TC_1", "review_comment": "Looks good"},
        {"id": "TC_4", "review_comment": "Clarify expected HTTP status code"},
    ]


@pytest.fixture
def sample_automation_results():
    return [
        {"id": "TC_1", "automatable": "Yes", "recommended_tool": "Selenium", "reason": "Stable UI flow"},
        {"id": "TC_2", "automatable": "Yes", "recommended_tool": "Selenium", "reason": "Deterministic negative check"},
        {"id": "TC_3", "automatable": "Yes", "recommended_tool": "Playwright", "reason": "Simple empty-state check"},
        {"id": "TC_4", "automatable": "No", "recommended_tool": "Manual", "reason": "Requires manual permission audit"},
        {"id": "TC_5", "automatable": "Yes", "recommended_tool": "JMeter", "reason": "Load/performance scenario"},
    ]


@pytest.fixture
def sample_estimation():
    return {"total_test_cases": 5, "test_cases_per_day": 20, "estimated_days": 1}


# ---------------- CREWAI MOCKING ----------------

@pytest.fixture
def make_task_output():
    """Factory: build a crewai TaskOutput by hand, e.g. for guardrail unit tests."""
    def _make(name, raw, agent_role="mock-agent"):
        return TaskOutput(description=f"description for {name}", name=name, raw=raw, agent=agent_role)
    return _make


@pytest.fixture
def make_stage_outputs():
    """
    Build a {task_name: raw_json_string} mapping for patch_crew_kickoff, filling in an
    empty array/object default for any stage not explicitly given.
    """
    def _encode(value, default):
        if value is None:
            return json.dumps(default)
        if isinstance(value, str):
            return value
        return json.dumps(value)

    def _make(analyze=None, generate=None, review=None, fix=None, automation=None, estimate=None):
        return {
            "analyze": _encode(analyze, []),
            "generate": _encode(generate, []),
            "review": _encode(review, []),
            "fix": _encode(fix, []),
            "automation": _encode(automation, []),
            "estimate": _encode(estimate, {}),
        }
    return _make


@pytest.fixture
def patch_crew_kickoff(monkeypatch):
    """
    Returns apply(outputs_by_task_name) which monkeypatches crewai.Crew.kickoff so that,
    instead of calling any real agent/LLM, it assigns a canned TaskOutput to each task
    (looked up by task.name) and fires the crew's task_callback — exercising all of
    run_pipeline's real orchestration/extraction/reporting code with zero API calls.

    The returned function also exposes `.captured_tasks`, the real crewai Task objects
    seen on the most recent kickoff (with their real, fully-rendered `.description`), so
    a test can assert on what was actually asked of a given stage (e.g. that a changed
    `testcases_per_day` value reached the estimation task's prompt).
    """
    state = {}

    def _apply(outputs_by_task_name):
        state.clear()
        state.update(outputs_by_task_name)
        _apply.captured_tasks = []

        def fake_kickoff(self):
            _apply.captured_tasks = list(self.tasks)
            for t in self.tasks:
                raw = state.get(t.name, "")
                t.output = TaskOutput(
                    description=t.description,
                    name=t.name,
                    raw=raw,
                    agent=getattr(t.agent, "role", "mock-agent"),
                )
                if getattr(self, "task_callback", None):
                    self.task_callback(t.output)
            return None

        monkeypatch.setattr(Crew, "kickoff", fake_kickoff)

    _apply.captured_tasks = []
    return _apply


@pytest.fixture
def failing_crew_kickoff(monkeypatch):
    """Monkeypatch crewai.Crew.kickoff to raise, simulating exhausted guardrail retries."""
    def _apply(error_message="Task failed guardrail validation after 3 retries."):
        def fake_kickoff(self):
            raise RuntimeError(error_message)
        monkeypatch.setattr(Crew, "kickoff", fake_kickoff)
    return _apply


# ---------------- SAMPLE FILES FOR load_requirements() ----------------

@pytest.fixture
def sample_txt_file(tmp_path, sample_requirements_text):
    path = tmp_path / "sample.txt"
    path.write_text(sample_requirements_text, encoding="utf-8")
    return str(path)


@pytest.fixture
def sample_docx_file(tmp_path, sample_requirements_text):
    path = tmp_path / "sample.docx"
    doc = DocxDocument()
    for line in sample_requirements_text.strip().splitlines():
        doc.add_paragraph(line)
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "REQ_EXTRA"
    table.rows[0].cells[1].text = "Requirement supplied via a table row"
    doc.save(str(path))
    return str(path)


@pytest.fixture
def sample_xlsx_file(tmp_path, sample_requirement_items):
    path = tmp_path / "sample.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.append(["Req ID", "Requirement", "Category", "Priority"])
    for r in sample_requirement_items:
        ws.append([r["req_id"], r["requirement"], r["category"], r["priority"]])
    wb.save(str(path))
    return str(path)


@pytest.fixture
def corrupted_docx_file(tmp_path):
    path = tmp_path / "corrupted.docx"
    path.write_bytes(b"this is not a real docx file")
    return str(path)
