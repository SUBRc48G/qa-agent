"""
Task 5: the test-case-generation stage (task2 / "generate") in isolation, via the
mocked crew. Covers schema validation, test-type coverage, and normalization with
automation lookup data.
"""

from openpyxl import load_workbook

import agent

REQUIRED_MANDATORY_TYPES = {"Positive", "Negative", "Edge", "Security", "Performance"}


def test_generated_test_cases_have_correct_schema(
    patch_crew_kickoff, make_stage_outputs, sample_test_cases, sample_requirements_text, tmp_output_dir
):
    patch_crew_kickoff(make_stage_outputs(generate=sample_test_cases))

    result = agent.run_pipeline(sample_requirements_text, output_dir=tmp_output_dir)

    assert result["counts"]["generated_test_cases"] == len(sample_test_cases)

    wb = load_workbook(result["files"]["test_cases"])
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    assert rows[0] == ("ID", "Req ID", "Scenario", "Type", "Steps", "Expected Result", "Priority", "Automatable")
    assert len(rows) == 1 + len(sample_test_cases)


def test_all_mandatory_test_types_are_represented_in_the_sample_fixture(sample_test_cases):
    # Guards the shared fixture itself: if it stops covering the mandatory types,
    # every test in this file that relies on "good" sample data becomes less useful.
    types_present = {tc["type"] for tc in sample_test_cases}
    assert REQUIRED_MANDATORY_TYPES.issubset(types_present)


def test_test_case_types_are_preserved_through_the_pipeline(
    patch_crew_kickoff, make_stage_outputs, sample_test_cases, sample_requirements_text, tmp_output_dir
):
    patch_crew_kickoff(make_stage_outputs(generate=sample_test_cases))
    result = agent.run_pipeline(sample_requirements_text, output_dir=tmp_output_dir)

    wb = load_workbook(result["files"]["test_cases"])
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))[1:]
    types_in_report = {row[3] for row in rows}

    expected_types = {tc["type"] for tc in sample_test_cases}
    assert types_in_report == expected_types
    assert REQUIRED_MANDATORY_TYPES.issubset(types_in_report)


def test_steps_are_rendered_as_a_numbered_list_in_the_report(
    patch_crew_kickoff, make_stage_outputs, sample_test_cases, sample_requirements_text, tmp_output_dir
):
    patch_crew_kickoff(make_stage_outputs(generate=sample_test_cases))
    result = agent.run_pipeline(sample_requirements_text, output_dir=tmp_output_dir)

    wb = load_workbook(result["files"]["test_cases"])
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))[1:]
    steps_by_id = {row[0]: row[4] for row in rows}

    assert steps_by_id["TC_1"] == "1. Log in as Read-only User\n2. Navigate to credentials list\n3. Verify list is displayed"


def test_normalization_applies_automation_lookup_to_draft_test_cases(
    patch_crew_kickoff, make_stage_outputs, sample_test_cases, sample_automation_results,
    sample_requirements_text, tmp_output_dir,
):
    # normalize_tc's automation_lookup is built from the automation stage and applied to
    # every test-case sheet keyed by id, including the draft one — verify that wiring.
    patch_crew_kickoff(make_stage_outputs(generate=sample_test_cases, automation=sample_automation_results))
    result = agent.run_pipeline(sample_requirements_text, output_dir=tmp_output_dir)

    wb = load_workbook(result["files"]["test_cases"])
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))[1:]
    automatable_by_id = {row[0]: row[7] for row in rows}

    assert automatable_by_id["TC_1"] == "Yes"
    assert automatable_by_id["TC_4"] == "No"


def test_test_case_with_unknown_id_gets_blank_automatable(
    patch_crew_kickoff, sample_requirements_text, tmp_output_dir
):
    orphan_test_case = [{
        "id": "TC_ORPHAN", "req_id": "REQ_1", "scenario": "Not covered by automation stage",
        "type": "Positive", "steps": ["Do a thing"], "expected_result": "It works", "priority": "Low",
    }]
    patch_crew_kickoff({
        "analyze": "[]", "generate": __import__("json").dumps(orphan_test_case),
        "review": "[]", "fix": "[]", "automation": "[]", "estimate": "{}",
    })

    result = agent.run_pipeline(sample_requirements_text, output_dir=tmp_output_dir)

    wb = load_workbook(result["files"]["test_cases"])
    ws = wb.active
    row = list(ws.iter_rows(values_only=True))[1]
    # openpyxl reads an empty-string cell back as None, not "".
    assert row[7] in ("", None)


def test_single_test_case_object_is_wrapped_in_a_list(
    patch_crew_kickoff, sample_requirements_text, tmp_output_dir
):
    import json

    single_tc = {
        "id": "TC_1", "req_id": "REQ_1", "scenario": "Only one", "type": "Positive",
        "steps": ["Step one"], "expected_result": "Works", "priority": "High",
    }
    patch_crew_kickoff({
        "analyze": "[]", "generate": json.dumps(single_tc),
        "review": "[]", "fix": "[]", "automation": "[]", "estimate": "{}",
    })

    result = agent.run_pipeline(sample_requirements_text, output_dir=tmp_output_dir)
    assert result["counts"]["generated_test_cases"] == 1
