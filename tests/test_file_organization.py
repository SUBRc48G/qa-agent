"""
Task 13: Output file organization and naming.
Verifies that files are placed in correct directories with proper naming conventions.
"""

import os
from pathlib import Path

from openpyxl import load_workbook

import agent


def test_all_output_files_are_placed_in_output_directory(
    patch_crew_kickoff, make_stage_outputs, sample_test_cases, sample_requirements_text, tmp_output_dir
):
    patch_crew_kickoff(make_stage_outputs(generate=sample_test_cases))
    result = agent.run_pipeline(sample_requirements_text, output_dir=tmp_output_dir)

    for key, file_path in result["files"].items():
        assert os.path.dirname(file_path) == tmp_output_dir
        assert os.path.exists(file_path)


def test_file_paths_are_absolute_or_resolvable(
    patch_crew_kickoff, make_stage_outputs, sample_test_cases, sample_requirements_text, tmp_output_dir
):
    patch_crew_kickoff(make_stage_outputs(generate=sample_test_cases))
    result = agent.run_pipeline(sample_requirements_text, output_dir=tmp_output_dir)

    for key, file_path in result["files"].items():
        p = Path(file_path)
        assert p.exists(), f"File {key} at {file_path} does not exist"


def test_excel_files_are_valid_and_readable(
    patch_crew_kickoff, make_stage_outputs, sample_test_cases, sample_requirements_text, tmp_output_dir
):
    patch_crew_kickoff(make_stage_outputs(generate=sample_test_cases))
    result = agent.run_pipeline(sample_requirements_text, output_dir=tmp_output_dir)

    for key, file_path in result["files"].items():
        if file_path.endswith(".xlsx"):
            try:
                wb = load_workbook(file_path)
                assert len(wb.sheetnames) > 0
            except Exception as e:
                raise AssertionError(f"Failed to read Excel file {key}: {e}")


def test_timestamped_files_are_named_with_datetime_suffix(
    patch_crew_kickoff, make_stage_outputs, sample_test_cases, sample_requirements_text, tmp_output_dir
):
    patch_crew_kickoff(make_stage_outputs(fix=sample_test_cases))
    result = agent.run_pipeline(sample_requirements_text, output_dir=tmp_output_dir)

    if "final" in result["files"]:
        final_filename = os.path.basename(result["files"]["final"])
        # Should have underscore for timestamp separation
        assert "_" in final_filename
        assert "final" in final_filename
