"""
Task 4: the requirement-analysis stage (task1 / "analyze") in isolation, driven through
the real run_pipeline() with crewai's Crew.kickoff mocked out (see patch_crew_kickoff in
conftest.py) so no real LLM call happens, but every line of run_pipeline's own logic
(JSON extraction, normalization, Excel writing) executes for real.
"""

import json

from openpyxl import load_workbook

import agent


def test_requirement_items_have_correct_schema(
    patch_crew_kickoff, make_stage_outputs, sample_requirement_items, sample_requirements_text, tmp_output_dir
):
    patch_crew_kickoff(make_stage_outputs(analyze=sample_requirement_items))

    result = agent.run_pipeline(sample_requirements_text, output_dir=tmp_output_dir)

    reqs = result["data"]["requirements"]
    assert len(reqs) == len(sample_requirement_items)
    for r in reqs:
        assert set(r.keys()) == {"req_id", "requirement", "category", "priority"}
        assert r["req_id"] != ""
        assert r["requirement"] != ""

    assert result["counts"]["requirements"] == len(sample_requirement_items)
    assert "requirements" in result["files"]


def test_requirements_excel_report_matches_input(
    patch_crew_kickoff, make_stage_outputs, sample_requirement_items, sample_requirements_text, tmp_output_dir
):
    patch_crew_kickoff(make_stage_outputs(analyze=sample_requirement_items))
    result = agent.run_pipeline(sample_requirements_text, output_dir=tmp_output_dir)

    wb = load_workbook(result["files"]["requirements"])
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))

    assert rows[0] == ("Req ID", "Requirement", "Category", "Priority")
    assert len(rows) == 1 + len(sample_requirement_items)
    for row, expected in zip(rows[1:], sample_requirement_items):
        assert row == (expected["req_id"], expected["requirement"], expected["category"], expected["priority"])


def test_missing_fields_are_normalized_to_empty_strings(
    patch_crew_kickoff, make_stage_outputs, sample_requirements_text, tmp_output_dir
):
    incomplete_items = [
        {"req_id": "REQ_1", "requirement": "Only has an id and text"},
        {"requirement": "Missing a req_id entirely", "category": "Functional", "priority": "Low"},
    ]
    patch_crew_kickoff(make_stage_outputs(analyze=incomplete_items))

    result = agent.run_pipeline(sample_requirements_text, output_dir=tmp_output_dir)

    reqs = result["data"]["requirements"]
    assert reqs[0]["category"] == ""
    assert reqs[0]["priority"] == ""
    assert reqs[1]["req_id"] == ""


def test_duplicate_req_ids_are_preserved_not_deduplicated(
    patch_crew_kickoff, make_stage_outputs, sample_requirements_text, tmp_output_dir
):
    duplicated_items = [
        {"req_id": "REQ_1", "requirement": "First version", "category": "Functional", "priority": "High"},
        {"req_id": "REQ_1", "requirement": "Second version with same id", "category": "Functional", "priority": "High"},
    ]
    patch_crew_kickoff(make_stage_outputs(analyze=duplicated_items))

    result = agent.run_pipeline(sample_requirements_text, output_dir=tmp_output_dir)

    reqs = result["data"]["requirements"]
    assert len(reqs) == 2
    assert [r["req_id"] for r in reqs] == ["REQ_1", "REQ_1"]
    assert result["counts"]["requirements"] == 2


def test_special_characters_survive_the_full_json_and_excel_round_trip(
    patch_crew_kickoff, make_stage_outputs, sample_requirements_text, tmp_output_dir
):
    tricky_items = [
        {
            "req_id": "REQ_1",
            "requirement": 'Support unicode (café, ünïcödé), quotes ("nested"), commas, and\nnewlines',
            "category": "Functional",
            "priority": "High",
        },
    ]
    patch_crew_kickoff(make_stage_outputs(analyze=tricky_items))

    result = agent.run_pipeline(sample_requirements_text, output_dir=tmp_output_dir)

    reqs = result["data"]["requirements"]
    assert "café" in reqs[0]["requirement"]
    assert "ünïcödé" in reqs[0]["requirement"]

    wb = load_workbook(result["files"]["requirements"])
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    assert "café" in rows[1][1]


def test_single_requirement_object_is_wrapped_in_a_list(
    patch_crew_kickoff, sample_requirements_text, tmp_output_dir
):
    # Some LLM responses return one bare JSON object instead of a one-item array —
    # run_pipeline must normalize that into a list rather than crash.
    single_item = {"req_id": "REQ_1", "requirement": "Only one requirement", "category": "Functional", "priority": "High"}
    patch_crew_kickoff({
        "analyze": json.dumps(single_item),
        "generate": "[]", "review": "[]", "fix": "[]", "automation": "[]", "estimate": "{}",
    })

    result = agent.run_pipeline(sample_requirements_text, output_dir=tmp_output_dir)

    assert result["counts"]["requirements"] == 1
    assert result["data"]["requirements"][0]["req_id"] == "REQ_1"
