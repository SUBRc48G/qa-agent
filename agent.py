import os
import re
import ast
import json
from collections import Counter
from dotenv import load_dotenv
from datetime import datetime
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font
from pypdf import PdfReader
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from crewai import Agent, Task, Crew, LLM
from rag_store import retrieve_context

# ---------------- UTILITIES ----------------

def safe(val):
    if isinstance(val, list):
        return "\n".join(safe(x) for x in val)
    if isinstance(val, dict):
        return json.dumps(val)
    if val is None:
        return ""
    return str(val)


def extract_json(text):
    """Robust JSON extractor"""
    if not text:
        return None

    text = str(text).strip()

    # direct parse
    try:
        return json.loads(text)
    except:
        pass

    # try markdown code blocks (prefer last block)
    code_blocks = re.findall(r"```(?:json)?\s*\n(.*?)\n```", text, re.DOTALL)
    for block in reversed(code_blocks):
        try:
            return json.loads(block.strip())
        except:
            pass

    # if text starts with bracket, validate it's properly formed
    if text and text[0] in '[{':
        bracket_type = '[' if text[0] == '[' else '{'
        close_bracket = ']' if bracket_type == '[' else '}'
        depth = 0

        for end_idx in range(len(text)):
            if text[end_idx] == bracket_type:
                depth += 1
            elif text[end_idx] == close_bracket:
                depth -= 1
                if depth == 0:
                    json_str = text[:end_idx+1]
                    # check remaining content after closing bracket
                    remaining = text[end_idx+1:].strip()

                    # if no remaining content, parse the top-level structure
                    if not remaining:
                        try:
                            return json.loads(json_str)
                        except:
                            return None

                    # if remaining has structural garbage (closing bracket), reject
                    if remaining and remaining[0] == '}':
                        return None

                    # if remaining has any content (prose or other brackets),
                    # use regex to find all JSON and return the last valid one
                    break

        # reached end without closing bracket = truncated JSON
        if text.count(bracket_type) > text.count(close_bracket):
            return None

    # try greedy regex first (handles cases with text between same-type brackets)
    matches = re.findall(r"\{.*\}|\[.*\]", text, re.DOTALL)

    # try to parse each match in reverse order (prefer the last/rightmost valid JSON)
    for m in reversed(matches):
        try:
            return json.loads(m)
        except:
            pass

    # if greedy regex produced exactly one match that failed to parse,
    # check if it's the "known limitation" (objects with text between, not just primitives)
    if len(matches) == 1:
        m = matches[0]
        # Check if it looks like multiple objects with text between them
        # (pattern: }...text...{  with the text containing colons/quotes suggesting JSON keys)
        if '}\n' in m and '{' in m and re.search(r'\}\s*\w+.*\{', m):
            # further check: if it contains ":" it's likely JSON with keys, indicating known limitation
            if ':' in m:
                return None

    # if greedy regex found no valid JSON, try bracket-matching for embedded structures
    # but avoid matching nested structures (only match top-level structures in the text)
    candidates = []
    i = 0

    while i < len(text):
        char = text[i]
        if char == '{':
            depth = 0
            for end_idx in range(i, len(text)):
                if text[end_idx] == '{':
                    depth += 1
                elif text[end_idx] == '}':
                    depth -= 1
                    if depth == 0:
                        try:
                            parsed = json.loads(text[i:end_idx+1])
                            candidates.append((i, parsed))
                        except:
                            pass
                        i = end_idx + 1
                        break
            else:
                i += 1
        elif char == '[':
            depth = 0
            for end_idx in range(i, len(text)):
                if text[end_idx] == '[':
                    depth += 1
                elif text[end_idx] == ']':
                    depth -= 1
                    if depth == 0:
                        try:
                            parsed = json.loads(text[i:end_idx+1])
                            candidates.append((i, parsed))
                        except:
                            pass
                        i = end_idx + 1
                        break
            else:
                i += 1
        else:
            i += 1

    # return the last valid JSON found (rightmost in the text)
    if candidates:
        return candidates[-1][1]

    return None


def _json_array_guardrail(task_output):
    """Crew guardrail: forces a retry if the task didn't return a parseable JSON array/object."""
    data = extract_json(task_output.raw)
    if data is None:
        return (False, "Your last response was not valid JSON. Return ONLY a valid JSON array, "
                        "with no markdown fences and no explanation before or after it.")
    if isinstance(data, dict):
        data = [data]
    if not isinstance(data, list):
        return (False, "Return a JSON array of objects, not a single value.")
    return (True, task_output.raw)


def _json_object_guardrail(task_output):
    """Crew guardrail: forces a retry if the task didn't return a parseable JSON object."""
    data = extract_json(task_output.raw)
    if not isinstance(data, dict):
        return (False, "Your last response was not a valid JSON object. Return ONLY a single valid "
                        "JSON object, with no markdown fences and no explanation before or after it.")
    return (True, task_output.raw)


def normalize_req(r):
    return {
        "req_id": safe(r.get("req_id")),
        "requirement": safe(r.get("requirement")),
        "category": safe(r.get("category")),
        "priority": safe(r.get("priority")),
    }


def format_steps(steps):
    """Render steps as a clean, numbered, newline-separated list for Excel display."""
    if isinstance(steps, list):
        lines = []
        for i, s in enumerate(steps, 1):
            s = str(s).strip()
            if re.match(r"^\d+[\.\)]", s):
                lines.append(s)
            else:
                lines.append(f"{i}. {s}")
        return "\n".join(lines)
    return safe(steps)


def normalize_tc(tc, automation_lookup=None):
    tc_id = safe(tc.get("id"))
    automatable = (automation_lookup or {}).get(tc_id, {}).get("automatable", "")
    return {
        "id": tc_id,
        "req_id": safe(tc.get("req_id")),
        "scenario": safe(tc.get("scenario")),
        "type": safe(tc.get("type")),
        "steps": format_steps(tc.get("steps")),
        "expected_result": safe(tc.get("expected_result")),
        "priority": safe(tc.get("priority")),
        "automatable": automatable,
    }


def normalize_automation(a, scenario_lookup):
    tc_id = safe(a.get("id"))
    return {
        "id": tc_id,
        "scenario": scenario_lookup.get(tc_id, ""),
        "automatable": safe(a.get("automatable")),
        "recommended_tool": safe(a.get("recommended_tool")),
        "reason": safe(a.get("reason")),
    }


def style_sheet(ws, col_widths, wrap_cols=None):
    """Set column widths, wrap long text, top-align cells, and grow row height to fit content."""
    for idx, width in enumerate(col_widths, start=1):
        ws.column_dimensions[ws.cell(row=1, column=idx).column_letter].width = width

    wrap_cols = set(wrap_cols) if wrap_cols is not None else set(range(1, len(col_widths) + 1))

    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")

    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, max_col=len(col_widths)):
        max_lines = 1
        for cell in row:
            if cell.column in wrap_cols:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
            if isinstance(cell.value, str) and "\n" in cell.value:
                max_lines = max(max_lines, cell.value.count("\n") + 1)
        ws.row_dimensions[row[0].row].height = max(15, min(15 * max_lines, 409))

    ws.freeze_panes = "A2"


def load_requirements(file_path):
    """Extract requirement text from an uploaded .txt, .pdf, .docx or .xlsx file."""
    ext = os.path.splitext(file_path)[1].lower()

    if ext == ".txt":
        with open(file_path, "r", encoding="utf-8") as f:
            return f.read()

    elif ext == ".pdf":
        reader = PdfReader(file_path)
        return "".join(page.extract_text() or "" for page in reader.pages)

    elif ext == ".docx":
        document = Document(file_path)
        parts = [p.text for p in document.paragraphs if p.text.strip()]

        for table in document.tables:
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells if c.text.strip()]
                if cells:
                    parts.append(" | ".join(cells))

        return "\n".join(parts)

    elif ext == ".xlsx":
        wb = load_workbook(file_path, data_only=True)
        parts = []

        for ws in wb.worksheets:
            for row in ws.iter_rows(values_only=True):
                cells = [str(c).strip() for c in row if c is not None and str(c).strip()]
                if cells:
                    parts.append(" | ".join(cells))

        return "\n".join(parts)

    else:
        raise ValueError("Unsupported file format. Supported: .txt, .pdf, .docx, .xlsx")


def load_requirements_auto():
    for file in ["requirements.pdf", "qa_requirements.txt"]:
        if os.path.exists(file):
            return load_requirements(file)
    raise FileNotFoundError("No requirement file found")


# ---------------- INIT ----------------

load_dotenv()

api_key = os.getenv("OPENAI_API_KEY")
if not api_key:
    raise ValueError("Missing OPENAI_API_KEY")

model_name = os.getenv("MODEL_NAME", "gpt-4o-mini")
default_testcases_per_day = int(os.getenv("TEST_CASES_PER_DAY", 20))

llm = LLM(model=model_name)


# ---------------- AGENTS ----------------
# Agents hold no per-run state, so they are safe to create once and reuse across pipeline runs.

analyst = Agent(role="Analyst", goal="Break requirements down into a structured, itemized list", backstory="QA expert", llm=llm, verbose=True)

qa_engineer = Agent(role="QA Engineer", goal="Generate the best possible, highest-coverage test cases", backstory="QA expert", llm=llm, verbose=True)

reviewer = Agent(role="Reviewer", goal="Review test cases", backstory="Senior QA", llm=llm, verbose=True)

fixer = Agent(role="Fixer", goal="Fix test cases", backstory="Senior QA", llm=llm, verbose=True)

automation_analyst = Agent(
    role="Automation Analyst",
    goal="Assess automation feasibility for each test case and recommend the right automation tool",
    backstory="Senior SDET with deep expertise across UI, API, mobile and performance automation frameworks",
    llm=llm,
    verbose=True
)

estimator = Agent(role="Estimator", goal="Estimate effort", backstory="QA lead", llm=llm, verbose=True)

automation_engineer = Agent(
    role="Automation Engineer",
    goal="Write clean, runnable Playwright pytest scripts that implement the given manual test cases",
    backstory="SDET who writes production-grade Playwright test automation in Python",
    llm=llm,
    verbose=True
)

strategy_writer = Agent(
    role="QA Lead",
    goal="Write clear, professional test strategy narrative tailored to the project",
    backstory="Senior QA Lead who has authored test strategy documents for many enterprise projects",
    llm=llm,
    verbose=True
)


# ---------------- PIPELINE ----------------

STAGE_LABELS = {
    "analyze": "✅ Requirements analyzed",
    "generate": "✅ Test cases generated",
    "review": "✅ Test cases reviewed",
    "fix": "✅ Test cases fixed",
    "automation": "✅ Automation feasibility assessed",
    "estimate": "✅ Effort estimated",
}


def run_pipeline(requirements, testcases_per_day=None, output_dir=".", progress_callback=None):
    """
    Run the full QA crew (analyze -> generate -> review -> fix -> automation -> estimate)
    against the given requirements text and write all Excel reports into output_dir.

    progress_callback, if given, is called with a short human-readable string as each
    stage completes (e.g. "✅ Test cases generated"), so a UI can show live progress.

    Returns a dict describing what was produced, so a caller (CLI or UI) can display
    a summary and offer the generated files for download. Any stage whose output could
    not be parsed as valid JSON (even after guardrail retries) is listed under "warnings"
    instead of failing the whole run silently.
    """
    testcases_per_day = testcases_per_day or default_testcases_per_day
    os.makedirs(output_dir, exist_ok=True)

    def notify(message):
        if progress_callback:
            progress_callback(message)

    # ---------------- TASKS (fresh per run, since they embed this run's requirements text) ----------------

    task1 = Task(
        description=f"""
Analyze the requirements below and break them down into a structured, itemized list.

RULES:
- Output ONLY a valid JSON array
- No markdown
- No explanation
- Split the requirements into distinct, atomic requirement items
- Each item must have:
  req_id, requirement, category, priority
  - req_id: e.g. "REQ_1", "REQ_2"
  - category: "Functional" or "Non-Functional"
  - priority: "High", "Medium", or "Low"

FORMAT EXAMPLE:
[
  {{
    "req_id": "REQ_1",
    "requirement": "User must be able to log in using email and password",
    "category": "Functional",
    "priority": "High"
  }}
]

Requirements:
{requirements}
""",
        expected_output="JSON array of requirement items",
        agent=analyst,
        name="analyze",
        guardrail=_json_array_guardrail,
    )

    task2 = Task(
        description=f"""
Using the itemized requirements from the analysis above, generate the BEST possible set of test cases
that gives maximum coverage of the requirements.

RULES:
- Output ONLY valid JSON array
- No markdown
- No explanation
- For every requirement, evaluate it against ALL of the test types below. Do NOT skip
  a type just because the requirement text doesn't explicitly mention it — infer
  applicability from what the requirement DOES, using the trigger patterns given.
  Only skip a type if none of its triggers apply at all.

  MANDATORY TYPES (apply the trigger test below to every requirement; if it matches,
  you MUST write that test case):
  - Positive (Functional): always applies — the requirement works with valid input/flow
  - Negative: always applies — invalid input, missing/incorrect data, disallowed action
  - Edge/Boundary: always applies — min/max values, empty/null, zero results, very large input
  - Security: applies whenever the requirement mentions a user role, permission, scope,
    license, credentials, or restricted action -> write a test that tries to bypass or
    exceed that restriction (e.g. a Read-only role attempting a write/delete action,
    or an action attempted without the required license)
  - Performance/Load: applies whenever the requirement involves listing, searching,
    filtering, retrieving, deleting, or otherwise processing multiple/bulk records ->
    write a test verifying acceptable response time when the data set is very large
    (e.g. thousands of records) or under concurrent requests
  - Data Validation: applies whenever the requirement involves structured data fields
    (IDs, names, dates, formats) -> write a test checking format/type/length consistency

  CONDITIONAL TYPES (include only if the requirement genuinely involves this):
  - Usability/UX: only if the requirement describes a user-facing UI/workflow
  - Compatibility: only if multiple browsers, devices, OS, or API versions are implied
  - Integration: only if the requirement explicitly involves another module/system/API
  - Regression: only if the requirement modifies/replaces existing behavior
  - Accessibility: only if the requirement describes a user-facing UI screen

- Before finalizing, double check: does the final list contain at least one Security
  test and at least one Performance test for every requirement that has a role/permission
  aspect or a list/search/bulk aspect respectively? If not, add them.
- Prioritize high-value, realistic scenarios; do not pad with irrelevant conditional types
- Every applicable type from the MANDATORY list MUST be its own separate test case object.
  Never combine two different types (e.g. a Positive check and a Security check) into a
  single test case just to keep the list short — each gets its own id, scenario and steps.
- "Do not create duplicate or redundant test cases" means do not repeat the exact same
  scenario twice — it does NOT mean skipping a mandatory type to save space
- Every test case must reference the req_id of the requirement it covers
- Each test case must have:
  id, req_id, scenario, type, steps, expected_result, priority
  - type: one of "Positive", "Negative", "Edge", "Security", "Usability",
    "Performance", "Compatibility", "Integration", "Data Validation",
    "Regression", "Accessibility"
  - steps: a JSON array of short, discrete step strings (one action per element,
    NOT a single paragraph), e.g. ["Open the login page", "Enter valid credentials",
    "Click Login"]. Do not number the steps yourself, that is added afterwards.

FORMAT EXAMPLE:
[
  {{
    "id": "TC_1",
    "req_id": "REQ_1",
    "scenario": "Login test",
    "type": "Positive",
    "steps": ["Open the login page", "Enter valid credentials", "Click Login"],
    "expected_result": "User logs in",
    "priority": "High"
  }}
]

Requirements:
{requirements}
""",
        expected_output="JSON array only",
        agent=qa_engineer,
        context=[task1],
        name="generate",
        guardrail=_json_array_guardrail,
    )

    task3 = Task(
        description="""
Review test cases and return ONLY JSON:

[
  {"id":"TC_1","review_comment":"..."}
]
""",
        expected_output="JSON only",
        agent=reviewer,
        context=[task2],
        name="review",
        guardrail=_json_array_guardrail,
    )

    task4 = Task(
        description="""
Fix the test cases based on the review comments.

RULES:
- Output ONLY valid JSON array
- No markdown, no explanation
- The output must contain EVERY test case id from the original list. Do not drop, merge,
  or silently remove a test case — even ones with no review comment must be carried
  through unchanged.
- Only remove a test case if its review comment explicitly says it is invalid, unsafe,
  or an exact duplicate of another id
- For every other test case, apply the fix described in its review comment (if any) and
  keep its id, req_id and type

Return ONLY final JSON array.
""",
        expected_output="Final JSON only",
        agent=fixer,
        context=[task2, task3],
        name="fix",
        guardrail=_json_array_guardrail,
    )

    task5 = Task(
        description="""
Analyze the FINAL test cases and determine automation feasibility for each one.

RULES:
- Output ONLY valid JSON array
- No markdown
- No explanation
- For every test case decide:
  - automatable: "Yes" or "No"
  - recommended_tool: the most suitable tool/framework for that test case, e.g.
    Selenium/Playwright/Cypress for Web UI, Appium for Mobile, Postman/RestAssured/pytest+requests
    for API, JMeter/k6 for Performance, or "Manual" if it cannot/should not be automated
    (e.g. exploratory, usability, one-off, or requires human judgement)
  - reason: one short sentence justifying the decision

FORMAT EXAMPLE:
[
  {"id":"TC_1","automatable":"Yes","recommended_tool":"Selenium","reason":"Simple, stable UI login flow"}
]
""",
        expected_output="JSON array only",
        agent=automation_analyst,
        context=[task4],
        name="automation",
        guardrail=_json_array_guardrail,
    )

    task6 = Task(
        description=f"""
Estimate effort.

Return JSON:
{{
 "total_test_cases": number,
 "test_cases_per_day": {testcases_per_day},
 "estimated_days": number
}}
""",
        expected_output="Estimation JSON",
        agent=estimator,
        context=[task4],
        name="estimate",
        guardrail=_json_object_guardrail,
    )

    # ---------------- CREW ----------------

    def _on_task_complete(task_output):
        notify(STAGE_LABELS.get(task_output.name, f"✅ Completed: {task_output.name}"))

    crew = Crew(
        agents=[analyst, qa_engineer, reviewer, fixer, automation_analyst, estimator],
        tasks=[task1, task2, task3, task4, task5, task6],
        verbose=True,
        memory=False,
        task_callback=_on_task_complete,
    )

    print("🚀 Running Crew...")
    notify("🚀 Starting QA crew...")
    crew_error = None
    try:
        crew.kickoff()
        print("✅ Crew completed")
    except Exception as e:
        # A task can exhaust its guardrail retries (bad JSON from the LLM even after
        # several attempts) and raise instead of just returning bad output. Degrade
        # gracefully: keep whatever tasks did complete and surface this as a warning
        # rather than crashing the whole run.
        crew_error = str(e)
        print(f"❌ Crew execution failed: {crew_error}")
        notify(f"❌ Crew execution stopped early: {crew_error}")

    # ---------------- EXTRACT JSON ----------------

    requirement_items = extract_json(task1.output)
    test_cases = extract_json(task2.output)
    reviews = extract_json(task3.output)
    updated_cases = extract_json(task4.output)
    automation_results = extract_json(task5.output)
    estimation = extract_json(task6.output)

    if isinstance(requirement_items, dict):
        requirement_items = [requirement_items]
    if isinstance(test_cases, dict):
        test_cases = [test_cases]
    if isinstance(reviews, dict):
        reviews = [reviews]
    if isinstance(updated_cases, dict):
        updated_cases = [updated_cases]
    if isinstance(automation_results, dict):
        automation_results = [automation_results]

    print(f"🔢 Stage counts — requirements: {len(requirement_items or [])}, "
          f"generated test cases: {len(test_cases or [])}, "
          f"reviews: {len(reviews or [])}, "
          f"final test cases: {len(updated_cases or [])}, "
          f"automation results: {len(automation_results or [])}")

    warnings = []
    if crew_error:
        warnings.append(f"The QA crew stopped early: {crew_error}")
    if not requirement_items:
        warnings.append("Could not extract structured requirements from the analysis step — "
                         "the requirements report was not generated.")
    if not test_cases:
        warnings.append("Could not extract generated test cases — the draft test cases report was not generated.")
    if not updated_cases:
        warnings.append("Could not extract final test cases after review/fix — "
                         "the final test cases report was not generated.")
    if not automation_results:
        warnings.append("Could not extract automation feasibility results — "
                         "the automation feasibility report was not generated.")
    if not estimation:
        warnings.append("Could not extract effort estimation — the estimation report was not generated.")

    for w in warnings:
        print(f"⚠️ {w}")
        notify(f"⚠️ {w}")

    notify("📊 Writing Excel reports...")

    automation_lookup = {}
    if automation_results:
        for a in automation_results:
            automation_lookup[safe(a.get("id"))] = {
                "automatable": safe(a.get("automatable")),
                "recommended_tool": safe(a.get("recommended_tool")),
            }

    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    files = {}

    # ---------------- REQUIREMENTS EXCEL ----------------

    wb0 = Workbook()
    ws0 = wb0.active
    ws0.title = "Requirements"
    ws0.append(["Req ID", "Requirement", "Category", "Priority"])

    if requirement_items:
        for r in requirement_items:
            ws0.append(list(normalize_req(r).values()))

        style_sheet(ws0, col_widths=[10, 60, 16, 10])
        path = os.path.join(output_dir, "requirements_analysis.xlsx")
        wb0.save(path)
        files["requirements"] = path
        print(f"📊 {len(requirement_items)} requirement(s) saved to {path}")
    else:
        print("❌ No requirements extracted")

    # ---------------- TEST CASES EXCEL ----------------

    wb1 = Workbook()
    ws1 = wb1.active
    ws1.title = "Test Cases"
    ws1.append(["ID", "Req ID", "Scenario", "Type", "Steps", "Expected Result", "Priority", "Automatable"])

    if test_cases:
        for tc in test_cases:
            ws1.append(list(normalize_tc(tc, automation_lookup).values()))

        style_sheet(ws1, col_widths=[10, 10, 30, 14, 55, 35, 10, 14])
        path = os.path.join(output_dir, "test_cases.xlsx")
        wb1.save(path)
        files["test_cases"] = path
        print("📊 Test cases saved")
    else:
        print("❌ No test cases found")

    # ---------------- REVIEW EXCEL ----------------

    wb2 = Workbook()
    ws2 = wb2.active
    ws2.title = "Review"
    ws2.append(["ID", "Comment"])

    if reviews:
        for r in reviews:
            ws2.append([safe(r.get("id")), safe(r.get("review_comment"))])

        style_sheet(ws2, col_widths=[10, 80])
        path = os.path.join(output_dir, "review_comments.xlsx")
        wb2.save(path)
        files["review"] = path
        print("📊 Review saved")

    # ---------------- FINAL TEST CASES EXCEL ----------------

    wb3 = Workbook()
    ws3 = wb3.active
    ws3.title = "Final Cases"
    ws3.append(["ID", "Req ID", "Scenario", "Type", "Steps", "Expected Result", "Priority", "Automatable"])

    scenario_lookup = {}
    final_cases_normalized = []

    if updated_cases:
        for tc in updated_cases:
            norm = normalize_tc(tc, automation_lookup)
            ws3.append(list(norm.values()))
            scenario_lookup[norm["id"]] = norm["scenario"]
            final_cases_normalized.append(norm)

        style_sheet(ws3, col_widths=[10, 10, 30, 14, 55, 35, 10, 14])
        path = os.path.join(output_dir, f"final_{timestamp}.xlsx")
        wb3.save(path)
        files["final"] = path
        print("📊 Final test cases saved")

    # ---------------- AUTOMATION FEASIBILITY EXCEL ----------------

    automation_summary = None

    if automation_results:
        wb5 = Workbook()
        ws5 = wb5.active
        ws5.title = "Automation Feasibility"
        ws5.append(["ID", "Scenario", "Automatable", "Recommended Tool", "Reason"])

        normalized_automation = [normalize_automation(a, scenario_lookup) for a in automation_results]

        for a in normalized_automation:
            ws5.append(list(a.values()))

        style_sheet(ws5, col_widths=[10, 30, 14, 20, 45])

        total_cases = len(normalized_automation)
        automatable_count = sum(1 for a in normalized_automation if a["automatable"].strip().lower() == "yes")
        manual_count = total_cases - automatable_count
        automation_pct = round((automatable_count / total_cases) * 100, 1) if total_cases else 0

        tool_counts = Counter(
            a["recommended_tool"] for a in normalized_automation if a["automatable"].strip().lower() == "yes"
        )

        ws6 = wb5.create_sheet("Summary")
        ws6.append(["Metric", "Value"])
        ws6.append(["Total Test Cases", total_cases])
        ws6.append(["Automatable", automatable_count])
        ws6.append(["Manual Only", manual_count])
        ws6.append(["Automation Coverage %", automation_pct])
        ws6.append([])
        ws6.append(["Recommended Tool", "Test Case Count"])
        for tool, count in tool_counts.most_common():
            ws6.append([tool, count])

        style_sheet(ws6, col_widths=[25, 20])
        path = os.path.join(output_dir, f"automation_feasibility_{timestamp}.xlsx")
        wb5.save(path)
        files["automation"] = path

        automation_summary = {
            "total": total_cases,
            "automatable": automatable_count,
            "manual": manual_count,
            "coverage_pct": automation_pct,
            "tool_counts": dict(tool_counts),
        }

        print(f"📊 Automation feasibility saved: {automatable_count}/{total_cases} ({automation_pct}%) test cases can be automated")
        for tool, count in tool_counts.most_common():
            print(f"   - {tool}: {count} test case(s)")
    else:
        print("❌ No automation feasibility results found")

    # ---------------- ESTIMATION EXCEL ----------------

    if estimation:
        wb4 = Workbook()
        ws4 = wb4.active
        ws4.title = "Estimation"
        ws4.append(["Total TC", "Per Day", "Days"])
        ws4.append([
            estimation.get("total_test_cases", 0),
            estimation.get("test_cases_per_day", testcases_per_day),
            estimation.get("estimated_days", 0)
        ])

        style_sheet(ws4, col_widths=[15, 12, 10])
        path = os.path.join(output_dir, f"estimation_{timestamp}.xlsx")
        wb4.save(path)
        files["estimation"] = path
        print("📊 Estimation saved")

    notify("🎉 Pipeline completed")
    print("🎉 PIPELINE COMPLETED SUCCESSFULLY")

    return {
        "timestamp": timestamp,
        "output_dir": output_dir,
        "files": files,
        "warnings": warnings,
        "counts": {
            "requirements": len(requirement_items or []),
            "generated_test_cases": len(test_cases or []),
            "reviews": len(reviews or []),
            "final_test_cases": len(updated_cases or []),
        },
        "automation_summary": automation_summary,
        "estimation": estimation,
        "data": {
            "requirements": [normalize_req(r) for r in requirement_items] if requirement_items else [],
            "final_test_cases": final_cases_normalized,
        },
    }


# ---------------- TEST STRATEGY DOCUMENT ----------------

DEFAULT_ENTRY_CRITERIA = [
    "Requirements are finalized and reviewed",
    "Test environment and test data are available",
    "Test cases are reviewed and approved",
]

DEFAULT_EXIT_CRITERIA = [
    "All planned test cases have been executed",
    "No open Critical/High severity defects remain",
    "Test summary report is reviewed and approved",
]

DEFAULT_RISKS = [
    {"risk": "Ambiguous or changing requirements", "mitigation": "Maintain a requirement traceability matrix and re-baseline test cases on change"},
    {"risk": "Test environment instability", "mitigation": "Set up a dedicated, monitored test environment ahead of execution"},
]

ROLES_AND_RESPONSIBILITIES = [
    ("QA Lead", "Owns the test strategy, allocates resources, reviews plans/reports, and tracks quality metrics."),
    ("Test Engineer", "Designs, executes and maintains test cases, and logs/retests defects."),
    ("Automation Engineer", "Builds and maintains automation scripts for automatable test cases and integrates them into CI/CD."),
    ("Developer", "Fixes defects, supports root cause analysis, and performs unit/component testing."),
    ("Product Owner", "Clarifies requirements, prioritizes defects, and signs off on releases."),
]

DELIVERABLES = [
    "Test Strategy Document",
    "Test Plan",
    "Test Cases (draft and final, with automation feasibility)",
    "Automation Scripts",
    "Defect Reports",
    "Test Summary Report",
    "Requirement Traceability Matrix",
]

KPIS = [
    "Test Coverage % — proportion of requirements mapped to at least one executed test case",
    "Defect Density — number of defects per requirement/module",
    "Pass/Fail Rate — percentage of executed test cases passing vs failing",
    "Escaped Defects — defects found in production after release",
]

REPORTING_CADENCE = [
    "Daily test execution status shared during stand-up",
    "Weekly test summary report circulated to stakeholders",
    "Final Test Closure Report published at the end of the test cycle",
]


def _add_table(doc, headers, rows, style="Light List Accent 1"):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = style
    for cell, header in zip(table.rows[0].cells, headers):
        cell.text = header
        cell.paragraphs[0].runs[0].bold = True
    for row_data in rows:
        cells = table.add_row().cells
        for cell, value in zip(cells, row_data):
            cell.text = str(value)
    return table


def _add_bullets(doc, items):
    for item in items:
        doc.add_paragraph(str(item), style="List Bullet")


def generate_frontend_testcases(base_url, max_pages=10, output_dir=".", progress_callback=None):
    """
    Discover pages on a frontend application by traversing links from the base URL,
    then generate comprehensive test cases for each discovered page.

    Args:
        base_url: The starting URL of the frontend application (e.g. "http://localhost:8501")
        max_pages: Maximum number of pages to discover and analyze (default 10)
        output_dir: Directory to save the results
        progress_callback: Optional function to receive progress updates

    Returns:
        A dict containing:
        - discovered_pages: List of discovered page URLs
        - test_cases_by_page: Dict mapping page URLs to their test cases
        - files: Dict with paths to generated Excel files
        - summary: Dict with counts and statistics
    """
    import asyncio
    from playwright.async_api import async_playwright
    from urllib.parse import urljoin, urlparse
    from collections import defaultdict

    # Streamlit runs the script in a worker thread that has no event loop of its
    # own; some crewai/litellm internals call asyncio.get_event_loop() and blow up
    # with "There is no current event loop in thread" if one was never set here.
    try:
        asyncio.get_event_loop()
    except RuntimeError:
        asyncio.set_event_loop(asyncio.new_event_loop())

    os.makedirs(output_dir, exist_ok=True)

    warnings = []

    def notify(message):
        if progress_callback:
            progress_callback(message)

    # Step 1: Discover pages using Playwright
    notify("🔍 Discovering frontend pages...")

    async def discover_pages():
        pages = set()
        visited = set()
        queue = [base_url]
        base_domain = urlparse(base_url).netloc

        async with async_playwright() as p:
            browser = await p.chromium.launch(headless=True)
            context = await browser.new_context()

            while queue and len(pages) < max_pages:
                url = queue.pop(0)
                if url in visited:
                    continue

                visited.add(url)

                try:
                    page = await context.new_page()
                    await page.goto(url, wait_until="domcontentloaded", timeout=15000)
                    try:
                        await page.wait_for_load_state("networkidle", timeout=5000)
                    except Exception:
                        pass  # SPA may keep network busy (polling/websockets) — don't fail the whole page for this

                    # Extract all links from the page
                    links = await page.locator("a[href]").all()
                    for link in links[:30]:  # Limit links per page
                        try:
                            href = await link.get_attribute("href")
                            if href:
                                # Resolve relative URLs
                                full_url = urljoin(url, href)
                                # Only follow links within the same domain
                                if urlparse(full_url).netloc == base_domain:
                                    # Keep hash fragments for SPAs
                                    if full_url not in visited and len(pages) < max_pages:
                                        pages.add(full_url)
                                        queue.append(full_url)
                        except:
                            pass

                    pages.add(url)
                    await page.close()

                except Exception as e:
                    notify(f"⚠️ Error visiting {url}: {str(e)}")

            await browser.close()

        return sorted(list(pages))[:max_pages]

    # Run async discovery
    discovered_pages = asyncio.run(discover_pages())
    notify(f"✅ Discovered {len(discovered_pages)} pages")

    # Step 2: Generate page descriptions for AI analysis
    notify("📝 Analyzing each page...")

    async def extract_page_content():
        page_descriptions = {}
        async with async_playwright() as p:
            browser = await p.chromium.launch(headless=True)
            context = await browser.new_context()

            for i, page_url in enumerate(discovered_pages, 1):
                try:
                    page = await context.new_page()
                    await page.goto(page_url, wait_until="domcontentloaded", timeout=10000)
                    try:
                        await page.wait_for_load_state("networkidle", timeout=3000)
                    except Exception:
                        pass  # SPA may keep network busy — proceed with whatever rendered so far

                    # Hash-routed SPAs finish client-side routing after load; without this the
                    # default route's DOM gets captured instead of the requested route's.
                    await page.wait_for_timeout(1500)

                    title = await page.title()

                    # inner_text() returns rendered visible text only; text_content()
                    # would also pull in <script>/<style> source, which poisons the prompt.
                    body_text = ""
                    try:
                        body_text = await page.locator("body").inner_text(timeout=5000)
                        body_text = " ".join(body_text.split())[:1200]
                    except Exception:
                        pass

                    ui = await page.evaluate("""() => {
                        const vis = el => {
                            const r = el.getBoundingClientRect();
                            const s = window.getComputedStyle(el);
                            return r.width > 0 && r.height > 0 && s.visibility !== 'hidden' && s.display !== 'none';
                        };
                        const txt = el => (el.innerText || el.textContent || '').trim().replace(/\\s+/g, ' ').slice(0, 80);

                        const fields = [...document.querySelectorAll('input, textarea, select')]
                            .filter(vis).slice(0, 40).map(el => ({
                                tag: el.tagName.toLowerCase(),
                                type: el.type || null,
                                name: el.name || el.id || null,
                                formControlName: el.getAttribute('formcontrolname') || null,
                                label: (document.querySelector(`label[for="${el.id}"]`) || {}).innerText?.trim().slice(0, 60)
                                    || el.closest('label')?.innerText?.trim().slice(0, 60) || null,
                                value: (el.type === 'radio' || el.type === 'checkbox') ? el.value : null,
                                placeholder: el.placeholder || null,
                                required: el.required || el.getAttribute('aria-required') === 'true',
                                maxlength: el.maxLength > 0 ? el.maxLength : null,
                                pattern: el.pattern || null,
                                min: el.min || null,
                                max: el.max || null,
                                options: el.tagName.toLowerCase() === 'select'
                                    ? [...el.options].slice(0, 10).map(o => o.text.trim()) : null,
                                ariaLabel: el.getAttribute('aria-label') || null,
                            }));

                        const buttons = [...document.querySelectorAll('button, input[type=submit], input[type=button], [role=button]')]
                            .filter(vis).slice(0, 25).map(el => ({
                                text: txt(el) || el.value || el.getAttribute('aria-label') || null,
                                disabled: !!el.disabled,
                                type: el.type || null,
                                opensMenu: el.getAttribute('aria-haspopup') || null,
                            })).filter(b => b.text);

                        const links = [...document.querySelectorAll('a[href]')]
                            .filter(vis).slice(0, 30).map(el => ({
                                text: txt(el) || null,
                                href: el.getAttribute('href'),
                                opensNewTab: el.getAttribute('target') === '_blank' || null,
                            })).filter(l => l.text);

                        const headings = [...document.querySelectorAll('h1, h2, h3')]
                            .filter(vis).slice(0, 15).map(el => `${el.tagName}: ${txt(el)}`);

                        return {
                            headings,
                            fields,
                            buttons,
                            links,
                            formCount: document.querySelectorAll('form').length,
                            tableCount: document.querySelectorAll('table').length,
                            imagesMissingAlt: [...document.querySelectorAll('img')].filter(i => !i.alt).length,
                            imageCount: document.querySelectorAll('img').length,
                            hasFileUpload: !!document.querySelector('input[type=file]'),
                            hasPasswordField: !!document.querySelector('input[type=password]'),
                            hasPagination: /next|previous|page \\d/i.test(document.body.innerText || ''),
                        };
                    }""")

                    path = urlparse(page_url).path or "/"
                    page_name = path.strip("/").replace("/", " > ") or "Home"

                    lines = [f"URL: {page_url}", f"Page name: {page_name}", f"Title: {title}"]
                    if ui.get("headings"):
                        lines.append("Headings: " + "; ".join(ui["headings"]))
                    if ui.get("fields"):
                        lines.append(f"Input fields ({len(ui['fields'])}):")
                        for f in ui["fields"]:
                            attrs = [f"{k}={v}" for k, v in f.items() if v not in (None, False, [], "")]
                            lines.append("  - " + ", ".join(attrs))
                    if ui.get("buttons"):
                        lines.append("Buttons (no href - clicking these does NOT navigate anywhere on its own): " + "; ".join(
                            b["text"]
                            + (" (disabled)" if b["disabled"] else "")
                            + (" (opens a menu/submenu, does not navigate)" if b.get("opensMenu") else "")
                            for b in ui["buttons"]))
                    if ui.get("links"):
                        lines.append("Links (real navigable href): " + "; ".join(
                            f'{l["text"]} -> {l["href"]}' + (" (opens in a NEW TAB)" if l.get("opensNewTab") else "")
                            for l in ui["links"][:20]))
                    structural = []
                    if ui.get("formCount"):
                        structural.append(f'{ui["formCount"]} form(s)')
                    if ui.get("tableCount"):
                        structural.append(f'{ui["tableCount"]} table(s)')
                    if ui.get("imageCount"):
                        structural.append(f'{ui["imageCount"]} image(s), {ui["imagesMissingAlt"]} missing alt text')
                    if ui.get("hasFileUpload"):
                        structural.append("has file upload")
                    if ui.get("hasPasswordField"):
                        structural.append("has password field")
                    if ui.get("hasPagination"):
                        structural.append("appears to have pagination")
                    if structural:
                        lines.append("Structure: " + ", ".join(structural))
                    if body_text:
                        lines.append(f"Visible text: {body_text}")

                    page_descriptions[page_url] = "\n".join(lines)
                    notify(f"  [{i}/{len(discovered_pages)}] Analyzed: {page_name}")
                    await page.close()

                except Exception as e:
                    notify(f"⚠️ Error analyzing {page_url}: {str(e)}")
                    # Fallback: use basic description
                    path = urlparse(page_url).path or "/"
                    page_name = path.strip("/").replace("/", " > ") or "Home"
                    page_descriptions[page_url] = f"Page: {page_name} (URL: {page_url})"

            await browser.close()
        return page_descriptions

    page_descriptions = asyncio.run(extract_page_content())

    # Step 3: Generate test cases for each page using the QA engineer agent
    notify("🧪 Generating test cases...")

    test_cases_by_page = defaultdict(list)
    all_test_cases = []

    for i, (page_url, description) in enumerate(page_descriptions.items(), 1):
        try:
            task = Task(
                description=f"""
You are testing a real web page. Below is its actual inspected DOM structure — every input
field, button and link listed was really found on the page. Write exhaustive test cases
grounded in these SPECIFIC elements (use their real names/labels/button text), not generic ones.

=== PAGE UNDER TEST ===
{description}
=== END PAGE ===

COVERAGE RULES — work through each of these and write every case that applies:

A. PER INPUT FIELD listed above, write a separate case for each that applies:
   - valid input accepted
   - empty/blank when the field is required
   - wrong format (e.g. email without "@", non-numeric in a number field)
   - boundary length (at maxlength, one over maxlength, single character)
   - whitespace-only input, leading/trailing spaces
   - special characters and unicode
   - injection-style payloads (XSS `<script>`, SQL `' OR 1=1--`) are REJECTED/escaped
   - for password fields: minimum length, weak password, password vs confirm-password mismatch, masking is applied
   - for selects: each meaningful option, plus no-selection
   - for file uploads: allowed type, disallowed type, oversized file

B. PER BUTTON listed above:
   - the happy path it triggers
   - clicking it with the form empty/invalid
   - double-click / rapid repeat submission does not duplicate the action
   - disabled-state behaviour where relevant

C. PER LINK / NAVIGATION listed above:
   - it navigates to the correct destination
   - browser back button returns correctly
   - direct deep-link / page refresh preserves expected state

D. ALWAYS include, for the page as a whole:
   - page loads with all listed elements visible
   - Accessibility: keyboard-only tab order, focus visibility, screen-reader labels on the
     listed fields, images missing alt text (if any were reported above)
   - Usability/Responsive: mobile viewport, tablet, desktop
   - Security: accessing this page unauthenticated vs authenticated (if it looks protected)
   - Performance: load time under a slow 3G network profile
   - Negative/Error handling: backend returns 500, request times out, offline/no network
   - Compatibility: Chrome, Firefox, Safari/WebKit

RULES:
- Output ONLY a valid JSON array. No markdown fences, no prose before or after.
- Be thorough: a page with several input fields should produce 20-40 test cases.
  Do NOT stop early and do NOT summarise multiple checks into one case.
- Each case must be ONE atomic check with its own id. Never merge a positive check and a
  negative check into the same case.
- Reference real element names/labels from the page data above in scenario and steps.
- Fields per object: id, page_url, scenario, type, steps, expected_result, priority
  - type: one of "Positive", "Negative", "Edge", "FormValidation", "Security",
    "Usability", "Accessibility", "Navigation", "Performance", "Compatibility"
  - steps: JSON array of short discrete action strings, do not number them
  - priority: "High", "Medium" or "Low"

Example of the required shape (yours must be far more extensive and page-specific):
[
  {{"id":"TC_1","page_url":"{page_url}","scenario":"Submit registration with all valid details","type":"Positive","steps":["Open the page","Fill every required field with valid data","Click the submit button"],"expected_result":"Account is created and user is redirected","priority":"High"}},
  {{"id":"TC_2","page_url":"{page_url}","scenario":"Email field rejects address with no @ symbol","type":"FormValidation","steps":["Open the page","Enter 'userexample.com' in the email field","Submit the form"],"expected_result":"Inline validation error indicates an invalid email format and the form is not submitted","priority":"High"}}
]

Now produce the full, exhaustive JSON array for: {page_url}
""",
                expected_output="A large JSON array of atomic, page-specific test cases",
                agent=qa_engineer,
                guardrail=_json_array_guardrail,
            )

            crew = Crew(agents=[qa_engineer], tasks=[task], verbose=False, memory=False)
            crew.kickoff()

            page_test_cases = extract_json(task.output) or []
            if isinstance(page_test_cases, dict):
                page_test_cases = [page_test_cases]

            # Validate test cases
            valid_test_cases = []
            for tc in page_test_cases:
                if isinstance(tc, dict) and all(k in tc for k in ["id", "scenario", "type", "steps", "expected_result"]):
                    tc["page_url"] = page_url
                    valid_test_cases.append(tc)
                    all_test_cases.append(tc)
                    test_cases_by_page[page_url].append(tc)

            if not valid_test_cases:
                warnings.append(f"No valid test cases parsed for {page_url}. Raw LLM output: {safe(task.output)[:300]}")
            notify(f"  [{i}/{len(discovered_pages)}] Generated {len(valid_test_cases)} test cases for: {page_url}")

        except Exception as e:
            warnings.append(f"Error generating test cases for {page_url}: {e}")
            notify(f"⚠️ Error generating test cases for {page_url}: {str(e)}")

    # Step 4: Save results to Excel
    notify("💾 Saving results...")

    # Pages summary
    pages_wb = Workbook()
    pages_ws = pages_wb.active
    pages_ws.title = "Discovered Pages"
    pages_ws.append(["Page URL", "Test Cases Generated"])

    for page_url in discovered_pages:
        count = len(test_cases_by_page.get(page_url, []))
        pages_ws.append([page_url, count])

    style_sheet(pages_ws, [60, 20])
    pages_file = os.path.join(output_dir, f"frontend_pages_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx")
    pages_wb.save(pages_file)

    # Test cases by page
    testcases_wb = Workbook()
    testcases_ws = testcases_wb.active
    testcases_ws.title = "Test Cases"
    testcases_ws.append(["Page URL", "Test ID", "Scenario", "Type", "Steps", "Expected Result", "Priority"])

    for tc in all_test_cases:
        testcases_ws.append([
            tc.get("page_url", ""),
            tc.get("id", ""),
            tc.get("scenario", ""),
            tc.get("type", ""),
            format_steps(tc.get("steps", [])),
            tc.get("expected_result", ""),
            tc.get("priority", ""),
        ])

    style_sheet(testcases_ws, [40, 12, 30, 15, 40, 40, 12], wrap_cols={2, 3, 4, 5, 6})
    testcases_file = os.path.join(output_dir, f"frontend_testcases_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx")
    testcases_wb.save(testcases_file)

    notify("✅ Frontend testing complete!")

    return {
        "discovered_pages": discovered_pages,
        "test_cases_by_page": dict(test_cases_by_page),
        "page_details": page_descriptions,
        "base_url": base_url,
        "output_dir": output_dir,
        "files": {
            "pages": pages_file,
            "test_cases": testcases_file,
        },
        "summary": {
            "total_pages_discovered": len(discovered_pages),
            "total_test_cases_generated": len(all_test_cases),
            "test_cases_per_page": {url: len(test_cases_by_page[url]) for url in discovered_pages},
        },
        "warnings": warnings,
    }


def _extract_code(text, page_url=None):
    """Pull python source out of an LLM reply, stripping markdown fences if present."""
    text = safe(text).strip()
    blocks = re.findall(r"```(?:python)?\s*\n(.*?)```", text, re.DOTALL)
    if blocks:
        code = max(blocks, key=len).strip()
    else:
        code = text
    code = _force_sync_playwright(code)
    code = _fix_playwright_re_import(code)
    if page_url:
        code = _force_correct_goto_url(code, page_url)
    return code


def _force_correct_goto_url(code, page_url):
    """
    The `base_url` fixture holds the single seed URL the whole suite was discovered from;
    each per-page test file is written for a DIFFERENT specific page and must navigate to
    its own page_url, not the shared fixture. The model occasionally confuses the two
    (probably because "base_url" reads as if it means "this page's URL") and writes
    `page.goto(base_url)`, silently making every test in that file load the wrong page.
    """
    return re.sub(r"page\.goto\(\s*base_url\s*\)", f'page.goto("{page_url}")', code)


_SINGULAR_LOCATOR_METHODS = {
    "get_by_role", "get_by_label", "get_by_placeholder",
    "get_by_text", "get_by_alt_text", "get_by_title",
}


def _disambiguate_locators_live(code, browser, page_url):
    """
    Mechanically catches the bug class that kept slipping through review: a locator that
    looks unique in isolation (e.g. get_by_label("Male")) but actually matches more than
    one element on the real page (e.g. "Female" contains "male" as a substring). Loads the
    real page once and asks Playwright itself how many elements each locator resolves to,
    then deterministically disambiguates with exact=True or, failing that, .first.

    Skips any locator that's deliberately meant to match multiple elements (assigned to a
    variable later used as a `for` loop's iterable, or used directly as one, or a bare
    get_by_role(role) with no name= — e.g. `for x in page.get_by_role("spinbutton"):`)
    since forcing those to a single match would silently break the test's intent.
    """
    try:
        tree = ast.parse(code)
    except SyntaxError:
        return code, 0

    protected = set()
    candidates = {}

    for func_node in ast.walk(tree):
        if not isinstance(func_node, ast.FunctionDef):
            continue

        assigned_call_by_var = {}
        for node in ast.walk(func_node):
            if isinstance(node, ast.Assign) and isinstance(node.value, ast.Call):
                src = ast.get_source_segment(code, node.value)
                if src:
                    for t in node.targets:
                        if isinstance(t, ast.Name):
                            assigned_call_by_var[t.id] = src

        for node in ast.walk(func_node):
            if isinstance(node, ast.For):
                it = node.iter
                src = assigned_call_by_var.get(it.id) if isinstance(it, ast.Name) else ast.get_source_segment(code, it)
                if src:
                    protected.add(src)

        for node in ast.walk(func_node):
            if not isinstance(node, ast.Call):
                continue
            fn = node.func
            if not (isinstance(fn, ast.Attribute) and isinstance(fn.value, ast.Name) and fn.value.id == "page"):
                continue
            if fn.attr not in _SINGULAR_LOCATOR_METHODS:
                continue
            try:
                args = [ast.literal_eval(a) for a in node.args]
                kwargs = {kw.arg: ast.literal_eval(kw.value) for kw in node.keywords if kw.arg}
            except (ValueError, TypeError):
                continue  # non-literal argument (e.g. a variable) - can't safely re-evaluate live
            if fn.attr == "get_by_role" and "name" not in kwargs:
                continue  # role-only lookups are usually intentionally plural
            src = ast.get_source_segment(code, node)
            if src and src not in candidates:
                candidates[src] = (fn.attr, args, kwargs)

    candidates = {src: v for src, v in candidates.items() if src not in protected}
    if not candidates:
        return code, 0

    try:
        page = browser.new_page()
        page.goto(page_url, wait_until="domcontentloaded", timeout=10000)
        try:
            page.wait_for_load_state("networkidle", timeout=3000)
        except Exception:
            pass
        page.wait_for_timeout(1000)
    except Exception:
        return code, 0

    fixed = 0
    try:
        for call_text, (method, args, kwargs) in candidates.items():
            try:
                count = getattr(page, method)(*args, **kwargs).count()
            except Exception:
                continue
            if count <= 1:
                continue

            replacement = None
            if not kwargs.get("exact"):
                try:
                    exact_kwargs = {**kwargs, "exact": True}
                    if getattr(page, method)(*args, **exact_kwargs).count() == 1:
                        replacement = call_text[:-1] + ", exact=True)"
                except Exception:
                    pass
            if replacement is None and not call_text.endswith(".first"):
                replacement = call_text + ".first"

            if replacement:
                code = code.replace(call_text, replacement)
                fixed += 1
    finally:
        page.close()

    return code, fixed


def _fix_playwright_re_import(code):
    """
    The model sometimes hallucinates `re` as a playwright.sync_api export (it isn't -
    `re` is the stdlib module used for regex matchers like re.compile). Strip it from
    the playwright import line and add a real `import re` if one isn't already present.
    """
    match = re.search(r"^from playwright\.sync_api import (.+)$", code, re.MULTILINE)
    if not match:
        return code
    names = [n.strip() for n in match.group(1).split(",")]
    if "re" not in names:
        return code
    names.remove("re")
    new_line = "from playwright.sync_api import " + ", ".join(names)
    code = code[: match.start()] + new_line + code[match.end() :]
    if not re.search(r"^import re\s*$", code, re.MULTILINE):
        code = "import re\n" + code
    return code


def _force_sync_playwright(code):
    """
    Despite explicit sync-API instructions, the model sometimes writes async def/await
    (its more common Playwright training pattern). Playwright's sync and async APIs are
    method-for-method identical, so a straight token rewrite is a safe, deterministic fix
    rather than re-prompting and hoping for compliance.
    """
    if "async def" not in code and re.search(r"\bawait\b", code) is None:
        return code
    code = re.sub(r"\basync def\b", "def", code)
    code = re.sub(r"\bawait\s+", "", code)
    code = re.sub(r"\basync with\b", "with", code)
    code = code.replace("playwright.async_api", "playwright.sync_api")
    # leftover from async-style generation; meaningless (and unregistered) on sync tests
    code = re.sub(r"^[ \t]*@pytest\.mark\.asyncio\s*\n", "", code, flags=re.MULTILINE)
    code = code.replace("import pytest_asyncio\n", "")
    return code


def _slugify(url):
    from urllib.parse import urlparse
    parsed = urlparse(url)
    raw = (parsed.path or "") + (parsed.fragment or "")
    slug = re.sub(r"[^a-zA-Z0-9]+", "_", raw).strip("_").lower()
    return slug or "home"


CONFTEST_TEMPLATE = '''import os
import pytest
from playwright.sync_api import sync_playwright

BASE_URL = os.getenv("BASE_URL", "{base_url}")
HEADLESS = os.getenv("HEADLESS", "true").lower() != "false"
TIMEOUT = int(os.getenv("TIMEOUT", "30000"))


@pytest.fixture(scope="session")
def base_url():
    return BASE_URL


@pytest.fixture(scope="session")
def browser():
    with sync_playwright() as p:
        browser = p.chromium.launch(headless=HEADLESS)
        yield browser
        browser.close()


@pytest.fixture
def page(browser):
    context = browser.new_context(viewport={{"width": 1920, "height": 1080}})
    page = context.new_page()
    page.set_default_timeout(TIMEOUT)
    yield page
    context.close()
'''

PYTEST_INI = """[pytest]
addopts = -v --tb=short
testpaths = .
"""

SUITE_README = """# Generated Playwright Suite

Auto-generated from the QA Agent frontend analysis of `{base_url}`.

## One-time setup

```bash
pip install -r requirements.txt
playwright install chromium
```

## Run the tests

```bash
# all tests (headless)
pytest

# watch the browser while it runs
HEADLESS=false pytest          # Windows PowerShell: $env:HEADLESS="false"; pytest

# a single file
pytest {example_file}

# point the suite at a different environment
BASE_URL=https://staging.example.com pytest

# with an HTML report
pytest --html=report.html --self-contained-html
```

## Notes

These scripts are a starting point generated from the test cases. Selectors were taken from
the live DOM at generation time, but assertions for flows needing real credentials or
server state are marked with `pytest.skip` or `TODO` and need your input before they pass.
"""

SUITE_REQUIREMENTS = """pytest>=8.0.0
playwright>=1.40.0
pytest-playwright>=0.4.0
pytest-html>=4.1.0
"""


def generate_playwright_scripts(test_cases_by_page, base_url, page_details=None, output_dir=".", progress_callback=None):
    """
    Turn generated frontend test cases into a runnable local Playwright + pytest suite.

    Writes one test_*.py per page plus conftest.py, pytest.ini, requirements.txt and a
    README into <output_dir>/playwright_suite, and returns a dict describing what was
    written (suite_dir, files, counts, warnings).
    """
    suite_dir = os.path.join(output_dir, "playwright_suite")
    os.makedirs(suite_dir, exist_ok=True)

    page_details = page_details or {}
    warnings = []
    written_files = []

    def notify(message):
        if progress_callback:
            progress_callback(message)

    pages_with_cases = {url: tcs for url, tcs in test_cases_by_page.items() if tcs}
    if not pages_with_cases:
        raise ValueError("No test cases available to convert into Playwright scripts.")

    from playwright.sync_api import sync_playwright
    _pw_ctx = sync_playwright().start()
    _verification_browser = _pw_ctx.chromium.launch(headless=True)

    for i, (page_url, test_cases) in enumerate(pages_with_cases.items(), 1):
        slug = _slugify(page_url)
        filename = f"test_{slug}.py"

        case_lines = []
        for tc in test_cases:
            steps = tc.get("steps")
            steps_text = " -> ".join(steps) if isinstance(steps, list) else safe(steps)
            case_lines.append(
                f'- id={tc.get("id")} | type={tc.get("type")} | priority={tc.get("priority")}\n'
                f'  scenario: {tc.get("scenario")}\n'
                f'  steps: {steps_text}\n'
                f'  expected: {tc.get("expected_result")}'
            )
        cases_block = "\n".join(case_lines)

        task = Task(
            description=f"""
Write a single Python Playwright test file implementing the manual test cases below.

TARGET PAGE: {page_url}

ACTUAL PAGE STRUCTURE (selectors seen in the live DOM — use these, do not invent selectors):
{page_details.get(page_url, "(not available)")}

TEST CASES TO IMPLEMENT:
{cases_block}

REQUIREMENTS:
- Output ONLY Python source code. No markdown fences, no commentary.
- Use the Playwright SYNC API via pytest, with this exact fixture contract already provided
  by an existing conftest.py (do NOT redefine these, do NOT call sync_playwright yourself):
    * `page`      - a ready Playwright Page
    * `base_url`  - the application base URL string
- Import only: `import pytest`, `import re` (if needed), and
  `from playwright.sync_api import expect, Page`
- One test function per test case, named test_<lowercase_snake_case_of_scenario>, and put
  the test case id in the docstring.
- Navigate with `page.goto("{page_url}")`.
- Prefer resilient locators in this order: get_by_role, get_by_label, get_by_placeholder,
  then CSS with the real name/id attributes shown in the page structure above.
- If a field's `label` shown above is null/missing (common for `<select>` and grouped
  radio/checkbox inputs) but it has a `formControlName`, use
  `page.locator('[formcontrolname="..."]')` instead of guessing a label — a missing label
  usually means the visible on-page text is not actually wired to that element via `for`,
  so `get_by_label` would either match nothing (hang/timeout) or match the wrong element.
- NEVER use a bare `page.locator("text=...")` or `get_by_text(...)` with a short/generic
  string — Playwright runs in strict mode and raises "resolved to N elements" if more than
  one element contains that text (this is common: a link, a heading and a paragraph can
  all contain the same word). Use `get_by_role("link"/"button", name="...", exact=True)`
  with the exact visible text from the Buttons/Links list above, and if a locator could
  still match more than one element, add `.first` explicitly.
- The same strict-mode ambiguity applies to `get_by_label(...)` and `get_by_placeholder(...)`:
  by default they match on substring, not full text, so pairs like "Password"/"Confirm
  Password" or "Male"/"Female" (which contains "male") will both match the shorter string.
  Whenever two labels/placeholders on the page share a common substring, pass `exact=True`
  to every `get_by_label`/`get_by_placeholder` call involved, not just the ambiguous one.
- For `<select>` elements, always call `.select_option(label="...")` using the visible
  option text. Never pass a bare positional string — that matches against the option's
  `value` attribute, which frequently differs from its visible label (e.g. the value may be
  encoded as "1: Doctor" while the visible text is just "Doctor"), causing a silent
  no-match/timeout instead of a clear error.
- To verify a password field is masked, assert `to_have_attribute("type", "password")`.
  Never assert that the field's value differs from what you just filled — masking is a
  purely visual effect and never changes the underlying value, so that assertion can never
  pass.

- ASSERTIONS: use ONLY the methods listed here. Do NOT invent any other assertion method.
    expect(locator).to_be_visible() | to_be_hidden() | to_be_enabled() | to_be_disabled()
    expect(locator).to_be_editable() | to_be_empty() | to_be_checked() | to_be_focused()
    expect(locator).to_have_text(t) | to_contain_text(t) | to_have_value(v)
    expect(locator).to_have_count(n) | to_have_class(c)
    expect(locator).to_have_attribute(name, value)   <- BOTH arguments are required
    expect(page).to_have_url(u) | to_have_title(t)
  Any of the above can be negated with the `not_` form, e.g. expect(loc).not_to_be_visible().
  There is no assertion for scrolling, layout, colour, or "looks correct" — for those,
  assert element visibility at the relevant viewport size instead.

- NEVER assert on markup you were not shown in the page structure above. Specifically:
    * do NOT assert CSS classes such as 'is-invalid'/'error' unless they appear above
    * do NOT assert a redirect/success URL unless that URL appears above
- Only elements from the Links list (a real `href` you were shown) may be asserted to
  navigate to a URL. A plain `<button>` (from the Buttons list) has no href — clicking it
  may open a menu, submit a form, or run arbitrary JS, but you have NO evidence it changes
  the URL, so do NOT write `expect(page).to_have_url(...)` after clicking one. If a button
  is marked "(opens a menu/submenu, does not navigate)", assert that instead, e.g.
  `expect(button).to_have_attribute("aria-expanded", "true")` after clicking it. If a link
  is marked "(opens in a NEW TAB)", the current `page` will never navigate — use
  `with page.context.expect_page() as new_tab_info:` around the click, then assert on
  `new_tab_info.value.url`, not `page.url`.

- Single-page apps rewrite their own URL on load (e.g. "/client/" becomes
  "/client/#/auth/login"), so NEVER assert that the URL still equals the URL you navigated
  to. To prove a form was rejected / not submitted, assert that a form element is still
  visible, e.g. expect(page.get_by_placeholder("...")).to_be_visible().
  If you must assert a URL, match a fragment of the route with a regex:
  expect(page).to_have_url(re.compile(r"auth/login")).

- MANDATORY SKIPS — begin the function body with `pytest.skip("needs <specific thing>")`
  for any case whose expected outcome depends on something you cannot obtain here:
  a successful login, valid credentials, a registered/seeded account, a specific redirect
  target, email delivery, or a server-forced 500/timeout. Write the function, then skip it.
  Guessing a URL or an error string instead of skipping is a failure.

- For responsive/viewport cases use `page.set_viewport_size({{"width": w, "height": h}})`.
- For cross-browser cases, do not launch other browsers — the `page` fixture is Chromium;
  `pytest.skip("needs firefox/webkit runner")` instead.
- For keyboard/tab-order cases: the exact tab order is not knowable from the page data
  above, so do NOT assert that a specific field receives focus after a fixed number of
  Tab presses. Instead assert the general capability, e.g. press Tab repeatedly and
  assert `page.locator(":focus")` has count 1 (something became focused), or that a named
  element `to_be_focused()` only after you explicitly `.focus()` it yourself first.
- Group related tests in classes only if it improves readability.
- The file must be syntactically valid and import-safe.

Return the complete contents of {filename}.
""",
            expected_output="Complete, runnable Python Playwright pytest module source",
            agent=automation_engineer,
        )

        try:
            crew = Crew(agents=[automation_engineer], tasks=[task], verbose=False, memory=False)
            crew.kickoff()
            code = _extract_code(task.output, page_url)

            if not code:
                warnings.append(f"{filename}: model returned no code")
                continue

            try:
                compile(code, filename, "exec")
            except SyntaxError as e:
                warnings.append(f"{filename}: generated code had a syntax error ({e.msg} on line {e.lineno}); saved anyway for manual fixing")

            code, disambiguated = _disambiguate_locators_live(code, _verification_browser, page_url)
            if disambiguated:
                notify(f"  [{i}/{len(pages_with_cases)}] Auto-disambiguated {disambiguated} locator(s) in {filename} against the live page")

            risky = re.findall(r'(?:locator|get_by_text)\(\s*["\'](?:text=)?([^"\']{1,40})["\']\s*\)(?!\.first)', code)
            if risky:
                warnings.append(
                    f"{filename}: {len(risky)} locator(s) may match multiple elements and lack "
                    f"exact=True/.first (e.g. {risky[0]!r}) — review before trusting a pass/fail result"
                )

            path = os.path.join(suite_dir, filename)
            with open(path, "w", encoding="utf-8") as f:
                f.write(code.rstrip() + "\n")

            written_files.append(path)
            notify(f"  [{i}/{len(pages_with_cases)}] Wrote {filename} ({len(test_cases)} cases)")

        except Exception as e:
            warnings.append(f"{filename}: {e}")
            notify(f"⚠️ Failed to generate {filename}: {e}")

    _verification_browser.close()
    _pw_ctx.stop()

    if not written_files:
        raise RuntimeError("No Playwright scripts could be generated. " + " | ".join(warnings))

    # Scaffolding written deterministically rather than by the LLM.
    with open(os.path.join(suite_dir, "conftest.py"), "w", encoding="utf-8") as f:
        f.write(CONFTEST_TEMPLATE.format(base_url=base_url))
    with open(os.path.join(suite_dir, "pytest.ini"), "w", encoding="utf-8") as f:
        f.write(PYTEST_INI)
    with open(os.path.join(suite_dir, "requirements.txt"), "w", encoding="utf-8") as f:
        f.write(SUITE_REQUIREMENTS)
    with open(os.path.join(suite_dir, "README.md"), "w", encoding="utf-8") as f:
        f.write(SUITE_README.format(
            base_url=base_url,
            example_file=os.path.basename(written_files[0]),
        ))

    notify(f"✅ Playwright suite written to {suite_dir}")

    return {
        "suite_dir": os.path.abspath(suite_dir),
        "test_files": written_files,
        "total_files": len(written_files),
        "total_cases": sum(len(t) for t in pages_with_cases.values()),
        "warnings": warnings,
    }


def generate_test_strategy(requirement_items, final_test_cases, automation_summary, estimation, output_dir="."):
    """
    Build a Test Strategy .docx from already-computed pipeline data (requirements, final
    test cases, automation feasibility, estimation), with an LLM writing the narrative
    sections (objectives, approach, entry/exit criteria, risks) grounded in that data.

    Returns the path to the generated .docx.
    """
    os.makedirs(output_dir, exist_ok=True)

    type_counts = Counter(tc.get("type") for tc in final_test_cases if tc.get("type"))

    traceability = {}
    for tc in final_test_cases:
        traceability.setdefault(tc.get("req_id") or "Unmapped", []).append(tc.get("id", ""))

    req_summary = "\n".join(f'{r["req_id"]}: {r["requirement"]}' for r in requirement_items) or "No requirements available"
    type_summary = ", ".join(f"{t}: {c}" for t, c in type_counts.items()) or "No test cases available"

    if automation_summary:
        automation_line = (
            f'{automation_summary["automatable"]}/{automation_summary["total"]} '
            f'({automation_summary["coverage_pct"]}%) test cases automatable. '
            f'Tools: {", ".join(automation_summary["tool_counts"].keys()) or "n/a"}'
        )
    else:
        automation_line = "Automation feasibility not available"

    if estimation:
        estimation_line = (
            f'{estimation.get("total_test_cases", "n/a")} total test cases, '
            f'{estimation.get("test_cases_per_day", "n/a")} per day, '
            f'{estimation.get("estimated_days", "n/a")} estimated days'
        )
    else:
        estimation_line = "Estimation not available"

    task = Task(
        description=f"""
Write the narrative content for a Test Strategy document for this project. Base it on the
actual project data below — do not invent unrelated details.

Project requirements:
{req_summary}

Test case type coverage: {type_summary}
Automation feasibility: {automation_line}
Effort estimation: {estimation_line}

RULES:
- Output ONLY valid JSON
- No markdown
- No explanation
- Keep every text field concise and professional (2-5 sentences for paragraphs)
- Provide 3-5 entry_criteria, 3-5 exit_criteria, and 3-5 risks that are realistic for THIS
  project's requirements (e.g. if a requirement involves licensing, include a risk about
  license/environment availability during testing)

JSON shape:
{{
  "introduction": "1-2 sentence purpose of this document for this specific project",
  "objectives": "paragraph describing the testing objectives for this project",
  "test_approach_narrative": "paragraph explaining the test approach/levels used, referencing the type coverage above",
  "automation_narrative": "paragraph explaining the automation strategy, referencing the automation feasibility numbers above",
  "entry_criteria": ["...", "..."],
  "exit_criteria": ["...", "..."],
  "risks": [{{"risk": "...", "mitigation": "..."}}]
}}
""",
        expected_output="JSON object with the test strategy narrative sections",
        agent=strategy_writer,
    )

    crew = Crew(agents=[strategy_writer], tasks=[task], verbose=True, memory=False)
    print("🚀 Writing test strategy narrative...")
    try:
        crew.kickoff()
    except Exception as e:
        # Fall back to the default narrative text (defined below) rather than failing
        # the whole document just because the LLM narrative call didn't come back clean.
        print(f"⚠️ Test strategy narrative generation failed, using defaults: {e}")

    narrative = extract_json(task.output) or {}

    doc = Document()

    title = doc.add_heading("Test Strategy Document", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle = doc.add_paragraph(f"Generated on {datetime.now().strftime('%Y-%m-%d %H:%M')}")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_heading("1. Introduction", level=1)
    doc.add_paragraph(safe(narrative.get("introduction")) or
                       "This document defines the test strategy for the project based on the analyzed requirements.")

    doc.add_heading("2. Scope & Objectives", level=1)
    doc.add_paragraph("In Scope:")
    if requirement_items:
        _add_bullets(doc, [f'{r["req_id"]}: {r["requirement"]}' for r in requirement_items])
    else:
        _add_bullets(doc, ["No requirements were extracted."])
    doc.add_paragraph("Objectives:")
    doc.add_paragraph(safe(narrative.get("objectives")) or
                       "Validate that the system meets the documented requirements with adequate quality and risk coverage.")

    doc.add_heading("3. Test Approach", level=1)
    doc.add_paragraph(safe(narrative.get("test_approach_narrative")) or
                       "A mix of functional and non-functional testing types is applied based on requirement characteristics.")
    if type_counts:
        _add_table(doc, ["Test Type", "Test Case Count"],
                   sorted(type_counts.items(), key=lambda x: -x[1]))

    doc.add_heading("4. Automation Approach", level=1)
    doc.add_paragraph(safe(narrative.get("automation_narrative")) or
                       "Test cases suitable for automation are identified and mapped to the appropriate tool.")
    if automation_summary:
        doc.add_paragraph(
            f'Automatable: {automation_summary["automatable"]}/{automation_summary["total"]} '
            f'({automation_summary["coverage_pct"]}%)  |  Manual: {automation_summary["manual"]}'
        )
        if automation_summary["tool_counts"]:
            _add_table(doc, ["Recommended Tool", "Test Case Count"],
                       sorted(automation_summary["tool_counts"].items(), key=lambda x: -x[1]))
    else:
        doc.add_paragraph("Automation feasibility data not available.")

    doc.add_heading("5. Effort Estimation & Schedule", level=1)
    if estimation:
        _add_table(doc, ["Metric", "Value"], [
            ("Total Test Cases", estimation.get("total_test_cases", "-")),
            ("Test Cases / Day", estimation.get("test_cases_per_day", "-")),
            ("Estimated Days", estimation.get("estimated_days", "-")),
        ])
    else:
        doc.add_paragraph("Estimation data not available.")

    doc.add_heading("6. Roles & Responsibilities", level=1)
    _add_table(doc, ["Role", "Responsibilities"], ROLES_AND_RESPONSIBILITIES)

    doc.add_heading("7. Deliverables", level=1)
    _add_bullets(doc, DELIVERABLES)

    if traceability:
        doc.add_heading("Requirement Traceability Matrix", level=2)
        req_text_lookup = {r["req_id"]: r["requirement"] for r in requirement_items}
        _add_table(doc, ["Req ID", "Requirement", "Test Case IDs"], [
            (req_id, req_text_lookup.get(req_id, ""), ", ".join(tc_ids))
            for req_id, tc_ids in traceability.items()
        ], style="Table Grid")

    doc.add_heading("8. Entry & Exit Criteria", level=1)
    doc.add_paragraph("Entry Criteria:")
    _add_bullets(doc, narrative.get("entry_criteria") or DEFAULT_ENTRY_CRITERIA)
    doc.add_paragraph("Exit Criteria:")
    _add_bullets(doc, narrative.get("exit_criteria") or DEFAULT_EXIT_CRITERIA)

    doc.add_heading("9. Risks & Mitigation", level=1)
    risks = narrative.get("risks") or DEFAULT_RISKS
    _add_table(doc, ["Risk", "Mitigation"], [(safe(r.get("risk")), safe(r.get("mitigation"))) for r in risks])

    doc.add_heading("10. Metrics & Reporting", level=1)
    doc.add_paragraph("KPIs tracked:")
    _add_bullets(doc, KPIS)
    doc.add_paragraph("Reporting cadence & format:")
    _add_bullets(doc, REPORTING_CADENCE)
    tools_used = ", ".join((automation_summary or {}).get("tool_counts", {}).keys()) or "N/A"
    doc.add_paragraph(
        f"Tools: {tools_used} for automation; Excel-based reports generated by this QA pipeline "
        f"for requirements, test cases, review and estimation."
    )

    path = os.path.join(output_dir, f"test_strategy_{datetime.now().strftime('%Y%m%d_%H%M%S')}.docx")
    doc.save(path)
    print(f"📄 Test strategy document saved to {path}")

    return path


if __name__ == "__main__":
    print("🚀 Script started")
    requirements = load_requirements_auto()
    print("✅ Requirements loaded")
    run_pipeline(requirements)
