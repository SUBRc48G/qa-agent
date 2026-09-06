import os
import re
import json
from collections import Counter
from dotenv import load_dotenv
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from pypdf import PdfReader
from crewai import Agent, Task, Crew, LLM
from rag_store import retrieve_context

# ---------------- UTILITIES ----------------

def safe(val):
    if isinstance(val, list):
        return "\n".join(safe(x) for x in val)
    if isinstance(val, dict):
        return json.dumps(val)
    if val is None:
        return ""
    return str(val)


def extract_json(text):
    """Robust JSON extractor"""
    if not text:
        return None

    text = str(text).strip()

    # direct parse
    try:
        return json.loads(text)
    except:
        pass

    # fallback regex extraction
    matches = re.findall(r"\{.*\}|\[.*\]", text, re.DOTALL)

    for m in reversed(matches):
        try:
            return json.loads(m)
        except:
            continue

    return None


def normalize_req(r):
    return {
        "req_id": safe(r.get("req_id")),
        "requirement": safe(r.get("requirement")),
        "category": safe(r.get("category")),
        "priority": safe(r.get("priority")),
    }


def format_steps(steps):
    """Render steps as a clean, numbered, newline-separated list for Excel display."""
    if isinstance(steps, list):
        lines = []
        for i, s in enumerate(steps, 1):
            s = str(s).strip()
            if re.match(r"^\d+[\.\)]", s):
                lines.append(s)
            else:
                lines.append(f"{i}. {s}")
        return "\n".join(lines)
    return safe(steps)


def normalize_tc(tc, automation_lookup=None):
    tc_id = safe(tc.get("id"))
    automatable = (automation_lookup or {}).get(tc_id, {}).get("automatable", "")
    return {
        "id": tc_id,
        "req_id": safe(tc.get("req_id")),
        "scenario": safe(tc.get("scenario")),
        "type": safe(tc.get("type")),
        "steps": format_steps(tc.get("steps")),
        "expected_result": safe(tc.get("expected_result")),
        "priority": safe(tc.get("priority")),
        "automatable": automatable,
    }


def normalize_automation(a, scenario_lookup):
    tc_id = safe(a.get("id"))
    return {
        "id": tc_id,
        "scenario": scenario_lookup.get(tc_id, ""),
        "automatable": safe(a.get("automatable")),
        "recommended_tool": safe(a.get("recommended_tool")),
        "reason": safe(a.get("reason")),
    }


def style_sheet(ws, col_widths, wrap_cols=None):
    """Set column widths, wrap long text, top-align cells, and grow row height to fit content."""
    for idx, width in enumerate(col_widths, start=1):
        ws.column_dimensions[ws.cell(row=1, column=idx).column_letter].width = width

    wrap_cols = set(wrap_cols or range(1, len(col_widths) + 1))

    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")

    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, max_col=len(col_widths)):
        max_lines = 1
        for cell in row:
            if cell.column in wrap_cols:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
            if isinstance(cell.value, str) and "\n" in cell.value:
                max_lines = max(max_lines, cell.value.count("\n") + 1)
        ws.row_dimensions[row[0].row].height = max(15, min(15 * max_lines, 409))

    ws.freeze_panes = "A2"


def load_requirements(file_path):
    ext = os.path.splitext(file_path)[1].lower()

    if ext == ".txt":
        with open(file_path, "r", encoding="utf-8") as f:
            return f.read()

    elif ext == ".pdf":
        reader = PdfReader(file_path)
        return "".join(page.extract_text() or "" for page in reader.pages)

    else:
        raise ValueError("Only .txt and .pdf supported")


def load_requirements_auto():
    for file in ["requirements.pdf", "qa_requirements.txt"]:
        if os.path.exists(file):
            return load_requirements(file)
    raise FileNotFoundError("No requirement file found")


# ---------------- INIT ----------------

print("🚀 Script started")

load_dotenv()

api_key = os.getenv("OPENAI_API_KEY")
if not api_key:
    raise ValueError("Missing OPENAI_API_KEY")

model_name = os.getenv("MODEL_NAME", "gpt-4o-mini")
testcases_per_day = int(os.getenv("TEST_CASES_PER_DAY", 20))

llm = LLM(model=model_name)

requirements = load_requirements_auto()
print("✅ Requirements loaded")


# ---------------- AGENTS ----------------

analyst = Agent(role="Analyst", goal="Break requirements down into a structured, itemized list", backstory="QA expert", llm=llm, verbose=True)

qa_engineer = Agent(role="QA Engineer", goal="Generate the best possible, highest-coverage test cases", backstory="QA expert", llm=llm, verbose=True)

reviewer = Agent(role="Reviewer", goal="Review test cases", backstory="Senior QA", llm=llm, verbose=True)

fixer = Agent(role="Fixer", goal="Fix test cases", backstory="Senior QA", llm=llm, verbose=True)

automation_analyst = Agent(
    role="Automation Analyst",
    goal="Assess automation feasibility for each test case and recommend the right automation tool",
    backstory="Senior SDET with deep expertise across UI, API, mobile and performance automation frameworks",
    llm=llm,
    verbose=True
)

estimator = Agent(role="Estimator", goal="Estimate effort", backstory="QA lead", llm=llm, verbose=True)


# ---------------- TASKS ----------------

task1 = Task(
    description=f"""
Analyze the requirements below and break them down into a structured, itemized list.

RULES:
- Output ONLY a valid JSON array
- No markdown
- No explanation
- Split the requirements into distinct, atomic requirement items
- Each item must have:
  req_id, requirement, category, priority
  - req_id: e.g. "REQ_1", "REQ_2"
  - category: "Functional" or "Non-Functional"
  - priority: "High", "Medium", or "Low"

FORMAT EXAMPLE:
[
  {{
    "req_id": "REQ_1",
    "requirement": "User must be able to log in using email and password",
    "category": "Functional",
    "priority": "High"
  }}
]

Requirements:
{requirements}
""",
    expected_output="JSON array of requirement items",
    agent=analyst
)

task2 = Task(
    description=f"""
Using the itemized requirements from the analysis above, generate the BEST possible set of test cases
that gives maximum coverage of the requirements.

RULES:
- Output ONLY valid JSON array
- No markdown
- No explanation
- For every requirement, evaluate it against ALL of the test types below. Do NOT skip
  a type just because the requirement text doesn't explicitly mention it — infer
  applicability from what the requirement DOES, using the trigger patterns given.
  Only skip a type if none of its triggers apply at all.

  MANDATORY TYPES (apply the trigger test below to every requirement; if it matches,
  you MUST write that test case):
  - Positive (Functional): always applies — the requirement works with valid input/flow
  - Negative: always applies — invalid input, missing/incorrect data, disallowed action
  - Edge/Boundary: always applies — min/max values, empty/null, zero results, very large input
  - Security: applies whenever the requirement mentions a user role, permission, scope,
    license, credentials, or restricted action -> write a test that tries to bypass or
    exceed that restriction (e.g. a Read-only role attempting a write/delete action,
    or an action attempted without the required license)
  - Performance/Load: applies whenever the requirement involves listing, searching,
    filtering, retrieving, deleting, or otherwise processing multiple/bulk records ->
    write a test verifying acceptable response time when the data set is very large
    (e.g. thousands of records) or under concurrent requests
  - Data Validation: applies whenever the requirement involves structured data fields
    (IDs, names, dates, formats) -> write a test checking format/type/length consistency

  CONDITIONAL TYPES (include only if the requirement genuinely involves this):
  - Usability/UX: only if the requirement describes a user-facing UI/workflow
  - Compatibility: only if multiple browsers, devices, OS, or API versions are implied
  - Integration: only if the requirement explicitly involves another module/system/API
  - Regression: only if the requirement modifies/replaces existing behavior
  - Accessibility: only if the requirement describes a user-facing UI screen

- Before finalizing, double check: does the final list contain at least one Security
  test and at least one Performance test for every requirement that has a role/permission
  aspect or a list/search/bulk aspect respectively? If not, add them.
- Prioritize high-value, realistic scenarios; do not pad with irrelevant conditional types
- Every applicable type from the MANDATORY list MUST be its own separate test case object.
  Never combine two different types (e.g. a Positive check and a Security check) into a
  single test case just to keep the list short — each gets its own id, scenario and steps.
- "Do not create duplicate or redundant test cases" means do not repeat the exact same
  scenario twice — it does NOT mean skipping a mandatory type to save space
- Every test case must reference the req_id of the requirement it covers
- Each test case must have:
  id, req_id, scenario, type, steps, expected_result, priority
  - type: one of "Positive", "Negative", "Edge", "Security", "Usability",
    "Performance", "Compatibility", "Integration", "Data Validation",
    "Regression", "Accessibility"
  - steps: a JSON array of short, discrete step strings (one action per element,
    NOT a single paragraph), e.g. ["Open the login page", "Enter valid credentials",
    "Click Login"]. Do not number the steps yourself, that is added afterwards.

FORMAT EXAMPLE:
[
  {{
    "id": "TC_1",
    "req_id": "REQ_1",
    "scenario": "Login test",
    "type": "Positive",
    "steps": ["Open the login page", "Enter valid credentials", "Click Login"],
    "expected_result": "User logs in",
    "priority": "High"
  }}
]

Requirements:
{requirements}
""",
    expected_output="JSON array only",
    agent=qa_engineer,
    context=[task1]
)

task3 = Task(
    description="""
Review test cases and return ONLY JSON:

[
  {"id":"TC_1","review_comment":"..."}
]
""",
    expected_output="JSON only",
    agent=reviewer,
    context=[task2]
)

task4 = Task(
    description="""
Fix the test cases based on the review comments.

RULES:
- Output ONLY valid JSON array
- No markdown, no explanation
- The output must contain EVERY test case id from the original list. Do not drop, merge,
  or silently remove a test case — even ones with no review comment must be carried
  through unchanged.
- Only remove a test case if its review comment explicitly says it is invalid, unsafe,
  or an exact duplicate of another id
- For every other test case, apply the fix described in its review comment (if any) and
  keep its id, req_id and type

Return ONLY final JSON array.
""",
    expected_output="Final JSON only",
    agent=fixer,
    context=[task2, task3]
)

task5 = Task(
    description="""
Analyze the FINAL test cases and determine automation feasibility for each one.

RULES:
- Output ONLY valid JSON array
- No markdown
- No explanation
- For every test case decide:
  - automatable: "Yes" or "No"
  - recommended_tool: the most suitable tool/framework for that test case, e.g.
    Selenium/Playwright/Cypress for Web UI, Appium for Mobile, Postman/RestAssured/pytest+requests
    for API, JMeter/k6 for Performance, or "Manual" if it cannot/should not be automated
    (e.g. exploratory, usability, one-off, or requires human judgement)
  - reason: one short sentence justifying the decision

FORMAT EXAMPLE:
[
  {"id":"TC_1","automatable":"Yes","recommended_tool":"Selenium","reason":"Simple, stable UI login flow"}
]
""",
    expected_output="JSON array only",
    agent=automation_analyst,
    context=[task4]
)

task6 = Task(
    description=f"""
Estimate effort.

Return JSON:
{{
 "total_test_cases": number,
 "test_cases_per_day": {testcases_per_day},
 "estimated_days": number
}}
""",
    expected_output="Estimation JSON",
    agent=estimator,
    context=[task4]
)


# ---------------- CREW ----------------

crew = Crew(
    agents=[analyst, qa_engineer, reviewer, fixer, automation_analyst, estimator],
    tasks=[task1, task2, task3, task4, task5, task6],
    verbose=True,
    memory=False
)

print("🚀 Running Crew...")

crew.kickoff()

print("✅ Crew completed")


# ---------------- GET TASK OUTPUTS ----------------

task1_output = task1.output
task2_output = task2.output
task3_output = task3.output
task4_output = task4.output
task5_output = task5.output
task6_output = task6.output


print("🔍 TASK2 OUTPUT:", task2_output)


# ---------------- EXTRACT JSON ----------------

requirement_items = extract_json(task1_output)
test_cases = extract_json(task2_output)
reviews = extract_json(task3_output)
updated_cases = extract_json(task4_output)
automation_results = extract_json(task5_output)
estimation = extract_json(task6_output)

if isinstance(automation_results, dict):
    automation_results = [automation_results]

print(f"🔢 Stage counts — requirements: {len(requirement_items or [])}, "
      f"generated test cases: {len(test_cases or [])}, "
      f"reviews: {len(reviews or [])}, "
      f"final test cases: {len(updated_cases or [])}, "
      f"automation results: {len(automation_results or [])}")

automation_lookup = {}
if automation_results:
    for a in automation_results:
        automation_lookup[safe(a.get("id"))] = {
            "automatable": safe(a.get("automatable")),
            "recommended_tool": safe(a.get("recommended_tool")),
        }

timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")


# ---------------- TASK1 EXCEL: REQUIREMENTS ----------------

if isinstance(requirement_items, dict):
    requirement_items = [requirement_items]

wb0 = Workbook()
ws0 = wb0.active
ws0.title = "Requirements"

ws0.append(["Req ID", "Requirement", "Category", "Priority"])

if requirement_items:
    for r in requirement_items:
        ws0.append(list(normalize_req(r).values()))

    style_sheet(ws0, col_widths=[10, 60, 16, 10])
    wb0.save("requirements_analysis.xlsx")
    print(f"📊 {len(requirement_items)} requirement(s) saved to requirements_analysis.xlsx")
else:
    print("❌ No requirements extracted")


# ---------------- TASK2 EXCEL ----------------

if isinstance(test_cases, dict):
    test_cases = [test_cases]

wb1 = Workbook()
ws1 = wb1.active
ws1.title = "Test Cases"

ws1.append(["ID", "Req ID", "Scenario", "Type", "Steps", "Expected Result", "Priority", "Automatable"])

if test_cases:
    for tc in test_cases:
        ws1.append(list(normalize_tc(tc, automation_lookup).values()))

    style_sheet(ws1, col_widths=[10, 10, 30, 14, 55, 35, 10, 14])
    wb1.save("test_cases.xlsx")
    print("📊 Test cases saved")
else:
    print("❌ No test cases found")


# ---------------- TASK3 EXCEL ----------------

if isinstance(reviews, dict):
    reviews = [reviews]

wb2 = Workbook()
ws2 = wb2.active
ws2.title = "Review"

ws2.append(["ID", "Comment"])

if reviews:
    for r in reviews:
        ws2.append([safe(r.get("id")), safe(r.get("review_comment"))])

    style_sheet(ws2, col_widths=[10, 80])
    wb2.save("review_comments.xlsx")
    print("📊 Review saved")


# ---------------- TASK4 EXCEL: FINAL TEST CASES ----------------

if isinstance(updated_cases, dict):
    updated_cases = [updated_cases]

wb3 = Workbook()
ws3 = wb3.active
ws3.title = "Final Cases"

ws3.append(["ID", "Req ID", "Scenario", "Type", "Steps", "Expected Result", "Priority", "Automatable"])

scenario_lookup = {}

if updated_cases:
    for tc in updated_cases:
        norm = normalize_tc(tc, automation_lookup)
        ws3.append(list(norm.values()))
        scenario_lookup[norm["id"]] = norm["scenario"]

    style_sheet(ws3, col_widths=[10, 10, 30, 14, 55, 35, 10, 14])
    fname = f"final_{timestamp}.xlsx"
    wb3.save(fname)
    print("📊 Final test cases saved")


# ---------------- TASK5 EXCEL: AUTOMATION FEASIBILITY ----------------

if automation_results:
    wb5 = Workbook()
    ws5 = wb5.active
    ws5.title = "Automation Feasibility"

    ws5.append(["ID", "Scenario", "Automatable", "Recommended Tool", "Reason"])

    normalized_automation = [normalize_automation(a, scenario_lookup) for a in automation_results]

    for a in normalized_automation:
        ws5.append(list(a.values()))

    style_sheet(ws5, col_widths=[10, 30, 14, 20, 45])

    total_cases = len(normalized_automation)
    automatable_count = sum(1 for a in normalized_automation if a["automatable"].strip().lower() == "yes")
    manual_count = total_cases - automatable_count
    automation_pct = round((automatable_count / total_cases) * 100, 1) if total_cases else 0

    tool_counts = Counter(
        a["recommended_tool"] for a in normalized_automation if a["automatable"].strip().lower() == "yes"
    )

    ws6 = wb5.create_sheet("Summary")
    ws6.append(["Metric", "Value"])
    ws6.append(["Total Test Cases", total_cases])
    ws6.append(["Automatable", automatable_count])
    ws6.append(["Manual Only", manual_count])
    ws6.append(["Automation Coverage %", automation_pct])
    ws6.append([])
    ws6.append(["Recommended Tool", "Test Case Count"])
    for tool, count in tool_counts.most_common():
        ws6.append([tool, count])

    style_sheet(ws6, col_widths=[25, 20])
    fname = f"automation_feasibility_{timestamp}.xlsx"
    wb5.save(fname)

    print(f"📊 Automation feasibility saved: {automatable_count}/{total_cases} ({automation_pct}%) test cases can be automated")
    for tool, count in tool_counts.most_common():
        print(f"   - {tool}: {count} test case(s)")
else:
    print("❌ No automation feasibility results found")


# ---------------- TASK6 EXCEL: ESTIMATION ----------------

if estimation:
    wb4 = Workbook()
    ws4 = wb4.active
    ws4.title = "Estimation"

    ws4.append(["Total TC", "Per Day", "Days"])
    ws4.append([
        estimation.get("total_test_cases", 0),
        estimation.get("test_cases_per_day", testcases_per_day),
        estimation.get("estimated_days", 0)
    ])

    style_sheet(ws4, col_widths=[15, 12, 10])
    fname = f"estimation_{timestamp}.xlsx"
    wb4.save(fname)
    print("📊 Estimation saved")


print("🎉 PIPELINE COMPLETED SUCCESSFULLY")
